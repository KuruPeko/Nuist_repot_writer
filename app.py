# -*- coding: utf-8 -*-
"""
ReportWriter · 后端（南信大理工科实验报告）
=========================================
打开软件 → 填封面信息 → 丢入实验材料 → 自动交给 DSH 写实验报告。

只依赖 Python 标准库 + python-docx + Pillow（DSH 自带环境里都有）。
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import socket
import subprocess
import sys
import threading
import time
import traceback
import uuid
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote
import urllib.request

# ---------------------------------------------------------------- 基础路径
ROOT = Path(__file__).resolve().parent
WEB = ROOT / "web"
DSH_ACCOUNT_PATCH = ROOT / "assets" / "dsh-account.patch.yml"
DSH_AUTH_ENABLE_PATCH = ROOT / "assets" / "dsh-auth-enable.patch.yml"
# 任务目录默认放在用户目录下的纯 ASCII 路径。DSH 的 Windows 沙箱要给工作区
# 追加 ACL，当前用户必须对 work_root 有 WRITE_DAC/WRITE_OWNER；安装器会尝试
# 给自定义工作区补 FullControl，手工改 work_root 时也要确保目录有完全控制权。
TASKS = Path(os.environ.get("LOCALAPPDATA") or ROOT) / "ReportWriter" / "tasks"
CONFIG_PATH = ROOT / "config.json"
LOG_PATH = ROOT / "app.log"
_local_appdata = os.environ.get("LOCALAPPDATA")
COVER_HISTORY_PATH = ((Path(_local_appdata) / "ReportWriter") if _local_appdata else ROOT) / "cover_history.json"
COVER_HISTORY_LIMIT = 30
COVER_HISTORY_FIELDS = ("course", "title", "name", "cls", "sid", "teacher", "date",
                        "school", "school_font", "school_color")
_COVER_HISTORY_LOCK = threading.Lock()

SKILL_REPO_URL = "https://github.com/Yutai233/NUIST-stem-lab-report-skill"
SKILL_UPDATE_URL = ("https://codeload.github.com/Yutai233/NUIST-stem-lab-report-skill/"
                    "zip/refs/heads/main")
SKILL_UPDATE_STATE_PATH = ((Path(_local_appdata) / "ReportWriter") if _local_appdata else ROOT) / "skill_update_state.json"
_SKILL_UPDATE = {
    "checking": False, "available": False, "latest": "",
    "checked_at": "", "error": "", "repo_url": SKILL_REPO_URL,
}
_SKILL_UPDATE_LOCK = threading.Lock()

TASKS.mkdir(parents=True, exist_ok=True)

DEFAULT_CONFIG = {
    "dsh": r"D:\dsh\resources\runtime\cli\bin\dsh.cmd",
    "python": sys.executable,
    "vendor": "vendor",
    "skill_root": "~/.agents/skills/nuist-stem-lab-report",
    "work_root": "%LOCALAPPDATA%\\ReportWriter\\tasks",
    "timeout": 2700,
    "port": 8791,
    "open_browser": True,
    # 封面顶部的红色毛体校名（图片素材，留空则用文字兜底）
    "school_logo": "assets/校名.png",
    "school_logo_cm": 11.0,
    "school_text": "南京信息工程大学",
    # 校名默认样式：毛体=用图片素材；其余字体名=直接排文字
    "school_font": "毛体",
    "school_color": "#C41E23",
    # DSH 执行权限：workspace-write（默认，安全）| danger-full-access（允许 MATLAB 等外部程序）
    "dsh_permission_mode": "workspace-write",
    # 运行结束后删除本次 dsh headless 的会话，避免 DSH 侧栏堆出“未分组”
    "cleanup_dsh_session": True,
    # 启动时检查 nuist-stem-lab-report skill 是否有更新
    "skill_update_check": True,
}

CONFIG = dict(DEFAULT_CONFIG)


def find_dsh():
    """config 里没写 dsh（或写的路径不存在）时，按常见安装位置探测一遍。"""
    candidates = [
        Path(os.environ.get("LOCALAPPDATA") or "") / "Programs/dsh/resources/runtime/cli/bin/dsh.cmd",
        Path.home() / ".dsh/dsh/resources/runtime/cli/bin/dsh.cmd",
        Path("D:/dsh/resources/runtime/cli/bin/dsh.cmd"),
        Path("C:/dsh/resources/runtime/cli/bin/dsh.cmd"),
        Path(os.environ.get("ProgramFiles", "C:/Program Files")) / "dsh/resources/runtime/cli/bin/dsh.cmd",
    ]
    for c in candidates:
        try:
            if str(c) and c.exists():
                return str(c)
        except OSError:
            continue
    return ""


def load_config():
    global CONFIG, TASKS
    if CONFIG_PATH.exists():
        try:
            data = json.loads(CONFIG_PATH.read_text("utf-8-sig"))
            CONFIG.update({k: v for k, v in data.items() if v is not None})
        except Exception:
            log("config.json 解析失败，使用默认配置")
    else:
        save_config()
    CONFIG["work_root"] = os.path.expandvars(str(CONFIG["work_root"]))
    # 其余路径：支持 %ENV%、~，相对路径按本程序目录解析（换机器也能用）
    for key in ("python", "vendor", "skill_root", "school_logo"):
        v = str(CONFIG.get(key) or "")
        if not v:
            continue
        v = os.path.expandvars(os.path.expanduser(v))
        p = Path(v)
        if not p.is_absolute():
            p = ROOT / p
        CONFIG[key] = str(p)
    TASKS = Path(CONFIG["work_root"])
    TASKS.mkdir(parents=True, exist_ok=True)
    # 配的路径不存在（或干脆留空）就自动探测一下，别人 clone 下来不用先改两处
    dsh_now = str(CONFIG.get("dsh") or "").strip()
    if not dsh_now or not Path(dsh_now).exists():
        cand = shutil.which("dsh") or shutil.which("dsh.cmd") or find_dsh()
        if cand:
            CONFIG["dsh"] = cand
    if not Path(str(CONFIG["python"])).exists():
        CONFIG["python"] = sys.executable


def save_config():
    try:
        CONFIG_PATH.write_text(json.dumps(CONFIG, ensure_ascii=False, indent=2), "utf-8")
    except Exception:
        pass


def dsh_home() -> Path:
    """DSH 的数据目录；安装了 DSH 就会在这里保存登录凭据。"""
    return Path(os.environ.get("DSH_HOME") or (Path.home() / ".dsh"))


def _credentials_block(text: str, key: str) -> list[str]:
    """从 credentials.yaml 里取出某个记录键下面的缩进块（不依赖 PyYAML）。"""
    lines = text.splitlines()
    for i, line in enumerate(lines):
        m = re.match(r"^(\s*)" + re.escape(key) + r":\s*$", line)
        if not m:
            continue
        base = len(m.group(1))
        block = []
        for nxt in lines[i + 1:]:
            if not nxt.strip() or nxt.lstrip().startswith("#"):
                continue
            if len(nxt) - len(nxt.lstrip()) <= base:
                break
            block.append(nxt)
        return block
    return []


def dsh_auth_state() -> tuple[str, str]:
    """返回 (鉴权方式, 给日志看的说明)。

    api_key       : 启动环境里有 DEEPSEEK_API_KEY
    account       : DSH 桌面端已登录，且凭据可用于正式模型
    account_stale : 有账号记录，但 token 缺失 / 是 mock 凭据 / 来自本机开发环境
    missing       : 两者都没检测到
    """
    if str(os.environ.get("DEEPSEEK_API_KEY") or "").strip():
        return "api_key", "使用环境变量 DEEPSEEK_API_KEY"
    cred = dsh_home() / ".credentials.yaml"
    try:
        text = cred.read_text("utf-8-sig", errors="replace")
    except OSError:
        text = ""

    block = _credentials_block(text, "deepseek-account-platform/default")
    if block:
        token = issuer = ""
        for line in block:
            m = re.match(r"^\s*token:\s*(.+?)\s*$", line)
            if m:
                token = m.group(1).strip().strip("'\"")
            m = re.match(r"^\s*issuer:\s*(.+?)\s*$", line)
            if m:
                issuer = m.group(1).strip().strip("'\"")
        if not token:
            return "account_stale", "DSH 账号记录缺少 token，请重新登录 DeepSeek Harness"
        if token.startswith("dsh_mock_"):
            return "account_stale", "DSH 当前是 mock/开发凭据，不能用于正式模型，请重新登录"
        if not issuer:
            return "account_stale", "DSH 账号记录缺少 issuer，请重新登录 DeepSeek Harness"
        host = (urlparse(issuer).hostname or "").lower()
        if host in ("localhost", "127.0.0.1", "::1", "[::1]"):
            return "account_stale", "DSH 账号凭据来自本机开发环境，不能用于正式模型，请重新登录"
        return "account", "使用 DSH 已登录账号"

    if "DEEPSEEK_API_KEY" in text:
        return "api_key", "使用 DSH 凭据库里的 DEEPSEEK_API_KEY"
    return "missing", "未检测到 DSH 登录凭据或 DEEPSEEK_API_KEY"


def dsh_launch_command(api_key: str = "", session_id: str = "") -> tuple[list[str], str]:
    """组装 dsh 命令：有账号凭据时，用补丁临时切到账号路由。

    不会修改 ~/.dsh 里的全局 profile。若用户自己在 headless profile 里配置过
    agent-default-model，则尊重用户配置，不叠加补丁。
    界面临时填写的 API Key 优先于账号补丁。
    """
    args = [CONFIG["dsh"], "headless"]
    profile_patch = dsh_home() / "profiles" / "headless" / "cordis.patch.yml"
    custom_route = False
    try:
        custom_route = (profile_patch.exists()
                        and "agent-default-model" in profile_patch.read_text("utf-8-sig", errors="replace"))
    except OSError:
        pass

    mode, note = dsh_auth_state()
    if str(api_key or "").strip():
        mode, note = "api_key", "使用本次填写的 DEEPSEEK_API_KEY"
    if custom_route:
        note = "使用 DSH headless profile 的自定义模型配置"
        if mode == "account" and DSH_AUTH_ENABLE_PATCH.exists():
            # 用户的路由可能指向 deepseek-account，但 headless 默认把
            # credentials/authorization 关掉了；只补启用，不覆盖 provider/model。
            args += ["--patch", str(DSH_AUTH_ENABLE_PATCH)]
            note += "（已启用账号鉴权插件）"
        elif mode in ("missing", "account_stale"):
            note += "（账号凭据不可用，仍可能报鉴权错误）"
    elif mode == "account" and DSH_ACCOUNT_PATCH.exists():
        args += ["--patch", str(DSH_ACCOUNT_PATCH)]
        note = "使用 DSH 已登录账号（临时 headless 路由补丁）"
    elif mode == "account":
        note = "检测到 DSH 登录凭据，但缺少 assets/dsh-account.patch.yml"
    elif mode == "missing":
        note = "未检测到 DSH 登录凭据或 DEEPSEEK_API_KEY"

    if session_id:
        args += ["--session-id", str(session_id)]
    if dsh_supports_json():
        args += ["--json", "-"]
    else:
        note += "（该 DSH 不支持 --json，无法统计 token）"
        args.append("-")
    return args, note


_DSH_JSON_CACHE = None


def dsh_supports_json() -> bool:
    """缓存检测当前 dsh 是否支持 headless --json（老版本没有这个选项）。"""
    global _DSH_JSON_CACHE
    if _DSH_JSON_CACHE is not None:
        return _DSH_JSON_CACHE
    dsh = str(CONFIG.get("dsh") or "")
    try:
        r = subprocess.run(
            [dsh, "headless", "--help"],
            capture_output=True, timeout=30,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        text = (r.stdout or b"").decode("utf-8", "replace") + \
               (r.stderr or b"").decode("utf-8", "replace")
        _DSH_JSON_CACHE = "--json" in text
    except Exception:
        # 探测失败时按支持处理；真遇到老版本，pump 仍会退回普通文本解析。
        _DSH_JSON_CACHE = True
    return _DSH_JSON_CACHE


def cleanup_dsh_session(session_id: str, timeout: float = 6.0) -> bool:
    """删除本次 dsh headless 创建的会话，避免它在 DSH 侧栏堆成“未分组”。

    只按 DSH 返回的 session id 精确匹配，且限定在 $DSH_HOME/sessions 下。
    DSH 可能等进程退出后才落盘，所以最多重试 timeout 秒。
    """
    sid = Path(str(session_id or "")).name
    if not sid or sid != str(session_id or ""):
        return False
    root = (dsh_home() / "sessions").resolve()
    if not root.is_dir():
        return False
    deadline = time.time() + max(0.0, timeout)
    while True:
        matches = []
        try:
            for p in root.rglob(sid):
                if p.name != sid:
                    continue
                rp = p.resolve()
                if root in rp.parents:
                    matches.append(rp)
        except OSError:
            pass
        if matches:
            ok = False
            for p in matches:
                try:
                    if p.is_dir():
                        shutil.rmtree(p)
                    else:
                        p.unlink()
                    ok = True
                except OSError:
                    log(f"清理 DSH 会话失败：{p}\n{traceback.format_exc()}")
            for parent in {p.parent for p in matches}:
                try:
                    if parent != root and parent.is_dir() and not any(parent.iterdir()):
                        parent.rmdir()
                except OSError:
                    pass
            return ok
        if time.time() >= deadline:
            return False
        time.sleep(0.5)


PERMISSION_MODES = ("read-only", "workspace-write", "danger-full-access")
PERMISSION_LABELS = {
    "read-only": "只读",
    "workspace-write": "标准（仅任务目录）",
    "danger-full-access": "完全访问（DSH 沙箱关闭）",
}


def permission_mode(opts: dict | None = None) -> str:
    """本次运行的 DSH 权限；默认 workspace-write。"""
    value = str((opts or {}).get("permission_mode")
                or CONFIG.get("dsh_permission_mode")
                or "workspace-write").strip()
    return value if value in PERMISSION_MODES else "workspace-write"


def log(msg: str):
    line = f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}\n"
    try:
        with LOG_PATH.open("a", encoding="utf-8") as f:
            f.write(line)
    except Exception:
        pass


def load_cover_history() -> list:
    """读取本地封面填写历史（只保留已知字段，最多 30 条）。"""
    try:
        if not COVER_HISTORY_PATH.exists():
            return []
        data = json.loads(COVER_HISTORY_PATH.read_text("utf-8-sig"))
        records = data.get("records") if isinstance(data, dict) else data
        if not isinstance(records, list):
            return []
        out = []
        for item in records:
            if not isinstance(item, dict):
                continue
            rec = {k: str(item.get(k) or "").strip() for k in COVER_HISTORY_FIELDS}
            if any(rec.values()):
                rec["updated"] = str(item.get("updated") or "")
                out.append(rec)
            if len(out) >= COVER_HISTORY_LIMIT:
                break
        return out
    except Exception:
        log("读取封面历史失败\n" + traceback.format_exc())
        return []


def save_cover_history(info: dict) -> list:
    """保存一份封面填写记录；相同内容只更新时间并移到最前。"""
    record = {k: str((info or {}).get(k) or "").strip() for k in COVER_HISTORY_FIELDS}
    if not any(record.values()):
        return load_cover_history()
    record["updated"] = datetime.now().isoformat(timespec="seconds")
    with _COVER_HISTORY_LOCK:
        records = load_cover_history()
        key = tuple(record[k] for k in COVER_HISTORY_FIELDS)
        records = [r for r in records
                   if tuple(r.get(k, "") for k in COVER_HISTORY_FIELDS) != key]
        records.insert(0, record)
        records = records[:COVER_HISTORY_LIMIT]
        try:
            COVER_HISTORY_PATH.parent.mkdir(parents=True, exist_ok=True)
            tmp = COVER_HISTORY_PATH.with_suffix(".tmp")
            tmp.write_text(json.dumps({"version": 1, "records": records},
                                      ensure_ascii=False, indent=2), "utf-8")
            tmp.replace(COVER_HISTORY_PATH)
        except Exception:
            log("保存封面历史失败\n" + traceback.format_exc())
        return records


def clear_cover_history():
    with _COVER_HISTORY_LOCK:
        try:
            COVER_HISTORY_PATH.unlink(missing_ok=True)
        except Exception:
            log("清除封面历史失败\n" + traceback.format_exc())


def local_skill_signature() -> str:
    """本地 skill 的轻量指纹：关键文件内容 + 所有文件路径/大小/mtime。"""
    root = Path(str(CONFIG.get("skill_root") or ""))
    if not (root / "SKILL.md").is_file():
        return ""
    h = hashlib.sha256()
    for rel in ("SKILL.md", "README.md",
                "references/course-registry.md",
                "references/general-stem-lab-prompt.md"):
        p = root / rel
        if p.is_file():
            h.update(rel.encode("utf-8"))
            h.update(p.read_bytes())
    try:
        for p in sorted(root.rglob("*"), key=lambda x: str(x).lower()):
            if not p.is_file() or p.name in ("Thumbs.db", "desktop.ini"):
                continue
            try:
                st = p.stat()
                rel = p.relative_to(root).as_posix()
                h.update(rel.encode("utf-8"))
                h.update(f"|{st.st_size}|{st.st_mtime_ns}\n".encode("utf-8"))
            except OSError:
                continue
    except OSError:
        pass
    return h.hexdigest()


def _read_skill_update_state() -> dict:
    try:
        if SKILL_UPDATE_STATE_PATH.exists():
            data = json.loads(SKILL_UPDATE_STATE_PATH.read_text("utf-8-sig"))
            return data if isinstance(data, dict) else {}
    except Exception:
        log("读取 skill 更新状态失败\n" + traceback.format_exc())
    return {}


def _write_skill_update_state(remote_etag: str, local_sig: str):
    try:
        SKILL_UPDATE_STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
        tmp = SKILL_UPDATE_STATE_PATH.with_suffix(".tmp")
        tmp.write_text(json.dumps({
            "remote_etag": remote_etag,
            "local_signature": local_sig,
            "checked_at": datetime.now().isoformat(timespec="seconds"),
        }, ensure_ascii=False, indent=2), "utf-8")
        tmp.replace(SKILL_UPDATE_STATE_PATH)
    except Exception:
        log("写入 skill 更新状态失败\n" + traceback.format_exc())


def fetch_skill_etag() -> str:
    """读取远端 ZIP 的 ETag；Python SSL 失败时回退到系统 curl。"""
    first_error = None
    try:
        req = urllib.request.Request(
            SKILL_UPDATE_URL, method="HEAD",
            headers={"User-Agent": "ReportWriter/1.0"},
        )
        with urllib.request.urlopen(req, timeout=12) as resp:
            etag = str(resp.headers.get("ETag") or "").strip()
        if etag:
            return etag
    except Exception as e:
        first_error = e

    curl = shutil.which("curl.exe") or shutil.which("curl")
    if curl:
        try:
            r = subprocess.run(
                [curl, "-s", "-I", "-L", "--max-time", "12", SKILL_UPDATE_URL],
                capture_output=True, timeout=20,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            text = (r.stdout or b"").decode("utf-8", "replace")
            m = re.search(r"(?im)^ETag:\s*(.+?)\s*$", text)
            if m:
                return m.group(1).strip()
        except Exception:
            pass
    raise RuntimeError(f"GitHub 未返回版本标识{('：' + str(first_error)) if first_error else ''}")


def check_skill_update_worker():
    with _SKILL_UPDATE_LOCK:
        _SKILL_UPDATE["checking"] = True
        _SKILL_UPDATE["error"] = ""
        _SKILL_UPDATE["available"] = False
    try:
        local_sig = local_skill_signature()
        if not local_sig:
            raise RuntimeError("未找到 nuist-stem-lab-report skill")
        state = _read_skill_update_state()
        etag = fetch_skill_etag()
        available = False
        baseline = str(state.get("remote_etag") or "")
        old_local = str(state.get("local_signature") or "")
        if not baseline or (old_local and old_local != local_sig):
            # 第一次检查，或本地 skill 已被手动更新：建立新基线
            _write_skill_update_state(etag, local_sig)
        elif etag != baseline:
            available = True
            # 不更新 remote_etag：在用户确认前，每次启动继续提示
            _write_skill_update_state(baseline, local_sig)
        else:
            _write_skill_update_state(baseline, local_sig)
        with _SKILL_UPDATE_LOCK:
            _SKILL_UPDATE.update({
                "checking": False,
                "available": available,
                "latest": etag if available else "",
                "checked_at": datetime.now().isoformat(timespec="seconds"),
                "error": "",
            })
        log("skill 更新检查：" + ("有更新" if available else "已是最新或首次建立基线"))
    except Exception as e:
        with _SKILL_UPDATE_LOCK:
            _SKILL_UPDATE.update({
                "checking": False, "available": False, "latest": "",
                "checked_at": datetime.now().isoformat(timespec="seconds"),
                "error": str(e),
            })
        log(f"skill 更新检查失败：{e}")


def start_skill_update_check():
    if not CONFIG.get("skill_update_check", True):
        return
    with _SKILL_UPDATE_LOCK:
        if _SKILL_UPDATE.get("checking"):
            return
        _SKILL_UPDATE["checking"] = True
    threading.Thread(target=check_skill_update_worker, daemon=True).start()


def ack_skill_update():
    with _SKILL_UPDATE_LOCK:
        latest = str(_SKILL_UPDATE.get("latest") or "")
        available = bool(_SKILL_UPDATE.get("available"))
    if not available or not latest:
        return
    local_sig = local_skill_signature()
    _write_skill_update_state(latest, local_sig)
    with _SKILL_UPDATE_LOCK:
        _SKILL_UPDATE["available"] = False
        _SKILL_UPDATE["latest"] = ""
        _SKILL_UPDATE["checked_at"] = datetime.now().isoformat(timespec="seconds")


# ---------------------------------------------------------------- 课程资源映射
# 与 skill 的 references/course-registry.md 保持一致
COURSES = [
    {
        "key": "单片机", "name": "单片机",
        "aliases": ["单片机", "微机原理", "mcu", "51"],
        "prompt": "单片机/单片机实验报告通用提示词.md",
        "template": "单片机/单片机实验模板.docx",
        "examples": "单片机/example",
    },
    {
        "key": "电路分析实验", "name": "电路分析实验",
        "aliases": ["电路分析", "电路实验", "电路原理", "电路基础"],
        "prompt": "电路分析实验/电路实验通用版提示词.md",
        "template": "电路分析实验/电路实验实验模板.docx",
        "examples": "电路分析实验/example",
    },
    {
        "key": "高频电子线路", "name": "高频电子线路",
        "aliases": ["高频"],
        "prompt": "高频电子线路/高频电子线路通用版提示词.md",
        "template": "高频电子线路/高频电子线路通用实验报告模板.docx",
        "examples": "",
    },
    {
        "key": "工程制图", "name": "工程制图",
        "aliases": ["工程制图", "工程图学", "机械制图", "cad"],
        "prompt": "工程制图/工图报告通用提示词.md",
        "template": "工程制图/工图模板.docx",
        "examples": "",
    },
    {
        "key": "模电实验", "name": "模电实验",
        "aliases": ["模电", "模拟电子", "模拟电路"],
        "prompt": "模电实验/模电实验通用版提示词.md",
        "template": "模电实验/模电实验模板.docx",
        "examples": "",
    },
    {
        "key": "数电实验", "name": "数电实验",
        "aliases": ["数电", "数字电子", "数字电路", "逻辑电路"],
        "prompt": "数电实验/数电实验报告通用提示词.md",
        "template": "数电实验/数电实验模板.docx",
        "examples": "数电实验/example",
    },
    {
        "key": "信息论基础", "name": "信息论基础",
        "aliases": ["信息论"],
        "prompt": "信息论基础/信息论基础版提示词.md",
        "template": "信息论基础/信息论基础实验报告基础模板.docx",
        "examples": "",
    },
]


def match_course(text: str):
    t = (text or "").strip().lower()
    if not t:
        return None
    for c in COURSES:
        for a in c["aliases"]:
            if a.lower() in t:
                return c
    return None


# ---------------------------------------------------------------- 封面规格
# 参考往期课程实验报告的封面版式：
# 标题 26pt 黑体加粗居中，个人信息 14pt 居中，日期 12pt 居中，页面 A4。
PAGE = {
    "w_cm": 21.0, "h_cm": 29.7,
    "top_cm": 2.54, "bottom_cm": 2.54,
    "left_cm": 3.17, "right_cm": 3.17,
    "header_cm": 1.27, "footer_cm": 1.27,
}

TITLE_PT, SUBTITLE_PT, INFO_PT, DATE_PT = 26, 16, 14, 12
EXPECTED_SECONDS = 540        # 一次完整写作大约多久（进度条时间估算用，只影响观感）
# 顶部先放红色毛体校名（图片），再是报告标题；spacer 由 150pt 让位给校名高度
SPACER_PT, TITLE_AFTER_PT, SUB_AFTER_PT, INFO_AFTER_PT, DATE_BEFORE_PT = 84, 48, 36, 14, 40
SCHOOL_AFTER_PT = 20          # 校名与报告标题之间的间距
SCHOOL_RED = (196, 30, 35)    # 兜底文字的颜色（与素材红一致）
# 校名可选字体：毛体=图片素材（无对应字体文件），其余为系统字体名
SCHOOL_FONTS = ("毛体", "华文行楷", "楷体", "仿宋", "黑体", "宋体", "微软雅黑", "华文新魏")
INFO_LABEL_EM = 5.0           # 「姓名：」这类标签占几个字宽（汉字宽=1em），冒号后补空格补齐
INFO_FIELD_EM = 10.0          # 下划线栏位宽（em）——四行一样长，右端自然对齐
TEXT_W_PT = (PAGE["w_cm"] - PAGE["left_cm"] - PAGE["right_cm"]) / 2.54 * 72


def school_logo_path() -> Path | None:
    """红色毛体校名素材；找不到就返回 None（走文字兜底）。"""
    rel = str(CONFIG.get("school_logo") or "").strip()
    if not rel:
        return None
    p = Path(rel)
    p = p if p.is_absolute() else (ROOT / p)
    return p if p.exists() else None


def parse_color(s) -> tuple | None:
    """'#C41E23' / 'c41e23' / '#c12' → (r, g, b)；解析不了返回 None。"""
    t = str(s or "").strip().lstrip("#")
    if re.fullmatch(r"[0-9a-fA-F]{6}", t):
        return tuple(int(t[i:i + 2], 16) for i in (0, 2, 4))
    if re.fullmatch(r"[0-9a-fA-F]{3}", t):
        return tuple(int(c * 2, 16) for c in t)
    return None


def school_style(info: dict) -> tuple[str, str, tuple]:
    """从表单信息里解析 (字体键, 校名文字, 颜色)，缺省回落到 config。"""
    info = info or {}
    font = str(info.get("school_font") or CONFIG.get("school_font") or "毛体").strip() or "毛体"
    text = str(info.get("school") or CONFIG.get("school_text") or "南京信息工程大学").strip()
    color = (parse_color(info.get("school_color"))
             or parse_color(CONFIG.get("school_color"))
             or SCHOOL_RED)
    return font, text or "南京信息工程大学", color


_LOGO_CACHE: dict = {}


def logo_png(path: Path, color: tuple) -> bytes:
    """把单色毛体校名重新上色，返回 PNG 字节（同一素材+颜色只算一次）。"""
    key = (str(path), tuple(color))
    if key in _LOGO_CACHE:
        return _LOGO_CACHE[key]
    from PIL import Image
    src = Image.open(str(path)).convert("RGBA")
    solid = Image.new("RGBA", src.size, tuple(color) + (255,))
    solid.putalpha(src.getchannel("A"))
    buf = io.BytesIO()
    solid.save(buf, format="PNG", optimize=True)
    data = buf.getvalue()
    _LOGO_CACHE[key] = data
    return data


def cover_title(course: str, title: str) -> str:
    c = (course or "").strip()
    if not c:
        return (title or "").strip() or "实验报告"
    if c.endswith("报告"):
        return c
    if c.endswith("实验"):
        return c + "报告"
    return c + "实验报告"


def norm_date(s: str) -> str:
    s = (s or "").strip()
    if not s:
        return datetime.now().strftime("%Y 年 %m 月 %d 日")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if m:
        return f"{int(m.group(1))} 年 {int(m.group(2))} 月 {int(m.group(3))} 日"
    return s


def _em_width(s: str) -> float:
    """粗估一行字占多少 em：汉字/全角 1，ASCII 0.5。用于对齐下划线栏位。"""
    return sum(1.0 if ord(ch) >= 0x2E80 else 0.5 for ch in s)


def _pad_to_em(s: str, target: float) -> str:
    """用全角/半角空格把 s 补到 target em 宽（用于四个信息行等长）。"""
    need = target - _em_width(s)
    if need <= 0:
        return s
    full = int(need)
    pad = "　" * full
    if need - full >= 0.5:
        pad += " "
    return s + pad


def cover_lines(info: dict):
    """返回封面排版块列表（docx 与预览共用同一份规格）。"""
    blocks = []
    blocks.append({"type": "spacer", "pt": SPACER_PT})
    font_key, school_text, school_color = school_style(info)
    logo = school_logo_path()
    if font_key == "毛体" and logo is not None:
        # 毛体没有可分发字体，用图片素材；按用户选的颜色重新上色
        blocks.append({"type": "logo", "path": str(logo),
                       "data": logo_png(logo, school_color),
                       "cm": float(CONFIG.get("school_logo_cm") or 11.0),
                       "after": SCHOOL_AFTER_PT})
    else:
        ea = "华文行楷" if font_key == "毛体" else font_key
        blocks.append({"type": "school", "pt": 24, "ea": ea, "bold": False,
                       "text": school_text,
                       "color": school_color, "after": SCHOOL_AFTER_PT})
    blocks.append({"type": "title", "pt": TITLE_PT, "ea": "黑体", "bold": True,
                   "text": cover_title(info.get("course", ""), info.get("title", "")),
                   "after": TITLE_AFTER_PT})
    sub = (info.get("title") or "").strip()
    if sub:
        blocks.append({"type": "subtitle", "pt": SUBTITLE_PT, "ea": "黑体", "bold": False,
                       "text": sub, "after": SUB_AFTER_PT})
    rows = [
        ("姓名", info.get("name", "")),
        ("班级", info.get("cls", "")),
        ("学号", info.get("sid", "")),
        ("任课教师", info.get("teacher", "")),
    ]
    # 四个信息行：标签占 LABEL_EM、下划线栏位占 FIELD_EM，行首用固定左缩进定位，
    # 于是四个标签的左边缘、四个下划线的起止点都严格对齐。
    indent_pt = (TEXT_W_PT - (INFO_LABEL_EM + INFO_FIELD_EM) * INFO_PT) / 2
    for label, value in rows:
        value = (value or "").strip()
        head = _pad_to_em(f"{label}：", INFO_LABEL_EM)   # 冒号后补空格，让下划线起点一致
        field = _pad_to_em(value, INFO_FIELD_EM)
        blocks.append({"type": "info", "pt": INFO_PT, "ea": "宋体", "bold": False,
                       "text": head + field,
                       "segs": [(head, False), (field, True)],
                       "indent_pt": indent_pt, "align": "left",
                       "after": INFO_AFTER_PT})
    blocks.append({"type": "date", "pt": DATE_PT, "ea": "宋体", "bold": False,
                   "text": f"实验日期：{norm_date(info.get('date', ''))}",
                   "before": DATE_BEFORE_PT})
    return blocks


# ---------------------------------------------------------------- 封面 DOCX
def build_cover_docx(info: dict, out_path: Path) -> Path:
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def set_font(run, ea="宋体", ascii_="Times New Roman"):
        run.font.name = ascii_
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), ascii_)
        rFonts.set(qn("w:hAnsi"), ascii_)
        rFonts.set(qn("w:eastAsia"), ea)
        rFonts.set(qn("w:cs"), ascii_)

    doc = docx.Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(PAGE["w_cm"]), Cm(PAGE["h_cm"])
    sec.top_margin, sec.bottom_margin = Cm(PAGE["top_cm"]), Cm(PAGE["bottom_cm"])
    sec.left_margin, sec.right_margin = Cm(PAGE["left_cm"]), Cm(PAGE["right_cm"])
    sec.header_distance, sec.footer_distance = Cm(PAGE["header_cm"]), Cm(PAGE["footer_cm"])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    nfmt = normal.paragraph_format
    nfmt.space_before = Pt(0)
    nfmt.space_after = Pt(0)
    nfmt.line_spacing = 1.0

    for block in cover_lines(info):
        if block["type"] == "spacer":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(block["pt"])
            p.paragraph_format.line_spacing = Pt(1)
            continue
        if block["type"] == "logo":
            # 毛体校名：以图片形式插入（毛体没有对应的可分发字体），颜色已在 data 里染好
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = Pt(0)
            pf.space_after = Pt(block.get("after", 0))
            src = io.BytesIO(block["data"]) if block.get("data") else block["path"]
            p.add_run().add_picture(src, width=Cm(block["cm"]))
            continue
        p = doc.add_paragraph()
        if block.get("align") == "left":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Pt(block.get("indent_pt") or 0)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_before = Pt(block.get("before", 0))
        pf.space_after = Pt(block.get("after", 0))
        # 「姓名：张三」这类信息行：标签正常、值那一栏带下划线
        for txt, underlined in (block.get("segs") or [(block.get("text", ""), False)]):
            if not txt:
                continue
            run = p.add_run(txt)
            run.font.size = Pt(block["pt"])
            run.bold = bool(block["bold"])
            if underlined:
                run.underline = True
            if block.get("color"):
                from docx.shared import RGBColor
                run.font.color.rgb = RGBColor(*block["color"])
            set_font(run, block["ea"])

    # 封面之后分节，正文从第 2 页开始
    new_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    new_sec.page_width, new_sec.page_height = Cm(PAGE["w_cm"]), Cm(PAGE["h_cm"])
    new_sec.top_margin, new_sec.bottom_margin = Cm(PAGE["top_cm"]), Cm(PAGE["bottom_cm"])
    new_sec.left_margin, new_sec.right_margin = Cm(PAGE["left_cm"]), Cm(PAGE["right_cm"])
    new_sec.header_distance, new_sec.footer_distance = Cm(PAGE["header_cm"]), Cm(PAGE["footer_cm"])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# ---------------------------------------------------------------- 封面预览 PNG
FONT_FILES = {
    "黑体": r"C:\Windows\Fonts\simhei.ttf",
    "宋体": r"C:\Windows\Fonts\simsun.ttc",
    "楷体": r"C:\Windows\Fonts\simkai.ttf",
    "仿宋": r"C:\Windows\Fonts\simfang.ttf",
    "微软雅黑": r"C:\Windows\Fonts\msyh.ttc",
    "华文行楷": r"C:\Windows\Fonts\STXINGKA.TTF",
    "华文新魏": r"C:\Windows\Fonts\STXINWEI.TTF",
}


def render_cover_png(info: dict, dpi: int = 150) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    scale = dpi / 72.0
    W, H = int(PAGE["w_cm"] / 2.54 * dpi), int(PAGE["h_cm"] / 2.54 * dpi)
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)

    cx = W // 2
    y = PAGE["top_cm"] / 2.54 * dpi  # 正文起点（pt→px 统一用 dpi 换算）

    def font_of(ea, pt):
        path = FONT_FILES.get(ea, FONT_FILES["宋体"])
        return ImageFont.truetype(path, max(8, int(round(pt * scale))))

    for block in cover_lines(info):
        if block["type"] == "spacer":
            y += block["pt"] * scale
            continue
        if block["type"] == "logo":
            from PIL import Image
            if block.get("data"):
                logo = Image.open(io.BytesIO(block["data"])).convert("RGBA")
            else:
                logo = Image.open(block["path"]).convert("RGBA")
            w = block["cm"] / 2.54 * dpi
            h = w * logo.height / logo.width
            logo = logo.resize((int(round(w)), int(round(h))), Image.LANCZOS)
            img.paste(logo, (int(round(cx - w / 2)), int(round(y))), logo)
            y += h + block.get("after", 0) * scale
            continue
        f = font_of(block["ea"], block["pt"])
        if block.get("before"):
            y += block["before"] * scale
        line_h = block["pt"] * 1.32 * scale
        fill = "black"
        if block.get("color"):
            fill = tuple(block["color"])
        segs = block.get("segs")
        if segs:
            # 分段绘制，值那一栏下面画横线（与 DOCX 的下划线对应）；带左缩进的块从固定 x 起笔
            widths = [draw.textlength(t, font=f) for t, _ in segs]
            total = sum(widths)
            if block.get("indent_pt") is not None:
                x = PAGE["left_cm"] / 2.54 * dpi + block["indent_pt"] * scale
            else:
                x = cx - total / 2
            base_y = y + line_h / 2
            asc = f.getmetrics()[0]
            ul_w = max(1, int(round(block["pt"] * scale / 16)))
            for (txt, underlined), w in zip(segs, widths):
                if txt:
                    draw.text((x, base_y), txt, font=f, fill=fill, anchor="lm")
                if underlined and w > 0:
                    uy = base_y + asc * 0.34
                    draw.line([(x, uy), (x + w, uy)], fill=fill, width=ul_w)
                x += w
        else:
            draw.text((cx, y + line_h / 2), block["text"], font=f, fill=fill, anchor="mm")
        y += line_h
        if block.get("after"):
            y += block["after"] * scale

    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


# ---------------------------------------------------------------- 工具函数
def safe_name(s: str, fallback="未命名", limit=40) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", (s or "").strip())
    s = re.sub(r"\s+", " ", s).strip(" .")
    if not s:
        return fallback
    return s[:limit]


def human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"


def read_text_any(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "utf-16"):
        try:
            return path.read_text(encoding=enc)
        except Exception:
            continue
    return path.read_bytes().decode("utf-8", "replace")


IMG_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
TEXT_EXT = {".md", ".markdown", ".txt", ".csv", ".json", ".py", ".c", ".h", ".cpp",
            ".java", ".m", ".v", ".hdl", ".log", ".err", ".cs", ".ps1", ".sh",
            ".html", ".htm", ".xml", ".yaml", ".yml"}
DOCX_EXT = {".docx", ".doc"}
PDF_EXT = {".pdf"}
TEMPLATE_EXT = {".docx", ".doc", ".dotx", ".pdf"}
TEMPLATE_MAX_BYTES = 30 * 1024 * 1024
MATERIAL_EXT = IMG_EXT | PDF_EXT | DOCX_EXT | TEXT_EXT
FOLDER_MAX_FILES = 200
FOLDER_MAX_FILE_BYTES = 80 * 1024 * 1024
FOLDER_MAX_TOTAL_BYTES = 300 * 1024 * 1024
SKIP_DIR_NAMES = {".git", ".svn", "node_modules", "__pycache__", ".venv", "venv",
                  ".idea", ".vscode", "System Volume Information"}

MD_IMG_RE = re.compile(r"!\[[^\]]*\]\(\s*([^)\s]+)")
HTML_IMG_RE = re.compile(r"""<img[^>]+src=["']([^"']+)["']""", re.I)


def kind_of(ext: str) -> str:
    ext = ext.lower()
    if ext in IMG_EXT:
        return "图片"
    if ext in PDF_EXT:
        return "PDF"
    if ext in DOCX_EXT:
        return "Word"
    if ext in TEXT_EXT:
        return "文本"
    return "其它"


# ---------------------------------------------------------------- 任务
def redact_text(text: str, values) -> str:
    """把个人信息原文替换成等长星号（日志、回答、落盘记录用）。"""
    for v in values or []:
        if v and v in text:
            text = text.replace(v, "*" * len(v))
    return text


class RunState:
    def __init__(self):
        self.state = "idle"          # idle | running | done | failed | canceled
        self.lines = []              # [seq, stream, text]
        self.lock = threading.Lock()
        self.proc = None
        self.started = None
        self.finished = None
        self.exit_code = None
        self.answer = ""
        self.prompt_path = ""
        self.pct = 0                 # 进度条走到过的最大值（只进不退）
        self.redact = []             # 日志脱敏用的个人信息原文
        self.last_output = None      # DSH 最近一次输出（判断是否真的开始干活）
        self.session_id = ""
        self.streamed_text = ""
        self.pause_requested = False
        self.cancel_requested = False
        self.pause_reason = ""
        self.usage = {
            "inputTokens": 0, "outputTokens": 0,
            "cacheReadTokens": 0, "cacheWriteTokens": 0,
            "totalTokens": 0, "steps": 0,
        }

    def add(self, stream: str, text: str):
        text = text.rstrip("\r\n")
        if not text:
            return
        if stream in ("stdout", "stderr"):
            self.last_output = time.time()
        if self.redact:
            text = redact_text(text, self.redact)
        with self.lock:
            self.lines.append([len(self.lines), stream, text])

    def add_usage(self, usage: dict):
        if not isinstance(usage, dict):
            return
        def num(key):
            try:
                return int(usage.get(key) or 0)
            except (TypeError, ValueError):
                return 0
        fields = {k: num(k) for k in
                  ("inputTokens", "outputTokens", "cacheReadTokens", "cacheWriteTokens")}
        total = num("totalTokens") or sum(fields.values())
        for k, v in fields.items():
            self.usage[k] += v
        self.usage["totalTokens"] += total
        self.usage["steps"] += 1

    def handle_event(self, ev: dict):
        """把 DSH headless --json 的 run event 转成日志、回答和 token 统计。"""
        typ = str(ev.get("type") or "")
        if typ == "session":
            sid = str(ev.get("sessionId") or "").strip()
            if sid:
                self.session_id = sid
                self.add("app", f"DSH session：{sid}")
        elif typ == "thinking":
            text = str(ev.get("text") or "").strip()
            if text:
                self.add("stderr", text)
        elif typ == "text":
            text = str(ev.get("text") or "")
            if text:
                if self.redact:
                    text = redact_text(text, self.redact)
                self.streamed_text += text
                self.answer = self.streamed_text
                self.add("stdout", text)
        elif typ == "final":
            text = str(ev.get("text") or "")
            if text:
                if self.redact:
                    text = redact_text(text, self.redact)
                self.answer = text
                if not self.streamed_text:
                    self.add("stdout", text)
        elif typ == "status":
            if ev.get("phase") == "step_end" and isinstance(ev.get("usage"), dict):
                usage = ev["usage"]
                self.add_usage(usage)
                def num(key):
                    try:
                        return int(usage.get(key) or 0)
                    except (TypeError, ValueError):
                        return 0
                parts = [f"输入 {num('inputTokens'):,}", f"输出 {num('outputTokens'):,}"]
                if num("cacheReadTokens"):
                    parts.append(f"缓存读 {num('cacheReadTokens'):,}")
                if num("cacheWriteTokens"):
                    parts.append(f"缓存写 {num('cacheWriteTokens'):,}")
                step = ev.get("step", "?")
                self.add("app", f"Token（第 {step} 步）：" + " · ".join(parts) +
                         f"；累计 {self.usage['totalTokens']:,}")
        elif typ == "tool_call":
            tool = str(ev.get("tool") or "tool")
            detail = json.dumps(ev.get("input"), ensure_ascii=False) if ev.get("input") is not None else ""
            if len(detail) > 800:
                detail = detail[:800] + "…"
            self.add("app", f"🔧 {tool} {detail}".rstrip())
        elif typ == "tool_result":
            result = ev.get("result")
            if not isinstance(result, str):
                result = json.dumps(result, ensure_ascii=False) if result is not None else ""
            if len(result) > 4000:
                result = result[:4000] + "…"
            stream = "error" if str(ev.get("status") or "").lower() in ("failed", "error") else "stdout"
            if result:
                self.add(stream, result)
        elif typ == "error":
            self.add("error", str(ev.get("message") or json.dumps(ev, ensure_ascii=False)))
        else:
            self.add("stdout", json.dumps(ev, ensure_ascii=False))

    def snapshot(self, start: int = 0):
        with self.lock:
            return self.lines[start:], len(self.lines)


class Task:
    def __init__(self, tid: str, path: Path, info: dict):
        self.id = tid
        self.path = path
        self.info = info
        self.opts = {}               # {outname, course_key, export_pdf, extra}
        self.materials = []          # [{name, path(rel), kind, size, text, note, pages, images}]
        self.template = None         # 教师提供的实验报告模板（可选）
        self.last_usage = {}         # 上一次运行的 token 统计（历史任务用）
        self.paused = None           # 暂停中的 DSH 会话（可继续）
        self.run = RunState()
        self.created = datetime.now().isoformat(timespec="seconds")

    # -- 目录
    @property
    def mat_dir(self) -> Path:
        return self.path / "材料"

    @property
    def txt_dir(self) -> Path:
        return self.mat_dir / "抽取文本"

    @property
    def img_dir(self) -> Path:
        return self.mat_dir / "附件图"

    @property
    def out_dir(self) -> Path:
        return self.path / "输出"

    @property
    def cover_path(self) -> Path:
        return self.path / "封面.docx"

    @property
    def template_dir(self) -> Path:
        return self.path / "模板"

    def init_dirs(self):
        for d in (self.mat_dir, self.txt_dir, self.img_dir, self.out_dir, self.template_dir):
            d.mkdir(parents=True, exist_ok=True)

    # -- 持久化
    def save(self):
        data = {
            "id": self.id, "info": self.info, "created": self.created,
            "opts": self.opts,
            "materials": self.materials,
            "template": self.template,
            "last_usage": self.last_usage,
            "paused": self.paused,
        }
        try:
            (self.path / "task.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), "utf-8")
        except Exception:
            log("task.json 写入失败\n" + traceback.format_exc())

    @classmethod
    def load(cls, path: Path) -> "Task | None":
        fp = path / "task.json"
        if not fp.exists():
            return None
        try:
            data = json.loads(fp.read_text("utf-8"))
        except Exception:
            return None
        t = cls(data.get("id") or path.name, path, data.get("info") or {})
        t.opts = data.get("opts") or {}
        t.materials = data.get("materials") or []
        t.template = data.get("template") or None
        t.last_usage = data.get("last_usage") or {}
        t.paused = data.get("paused") or None
        if isinstance(t.paused, dict):
            u = dict(t.run.usage)
            src = t.paused.get("usage") or {}
            for k in u:
                try:
                    u[k] = int(src.get(k) or 0)
                except (TypeError, ValueError):
                    pass
            t.run.usage = u
            t.run.state = "paused"
            t.run.session_id = str(t.paused.get("session_id") or "")
            t.run.pause_reason = str(t.paused.get("reason") or "")
            try:
                t.run.pct = int(t.paused.get("pct") or 0)
            except (TypeError, ValueError):
                t.run.pct = 0
            t.last_usage = dict(u)
        t.created = data.get("created") or ""
        return t

    def progress(self) -> dict:
        """进度：产物里程碑优先，写作期用时间估算兜住。

        100% 只有一个含义：实验报告已经完成——主 DOCX 已落盘，或 PDF 已导出。
        DOCX 还在写、DSH 还没退出时，最多停在 96%，不会提前报 100%。
        """
        st = self.run.state
        out = self.out_dir
        docx = pdf = False
        # 重跑时输出目录里可能还留着上一轮的文件；只把本轮开始后写出的文件算作产物
        started_ts = None
        if self.run.started:
            try:
                started_ts = datetime.fromisoformat(self.run.started).timestamp() - 2
            except ValueError:
                started_ts = None
        if out.exists():
            try:
                for f in out.iterdir():
                    # 跳过 Word 的 ~$ 临时文件和空文件，避免把半成品当成报告
                    if not f.is_file() or f.name.startswith("~$"):
                        continue
                    try:
                        stat = f.stat()
                        if stat.st_size <= 0:
                            continue
                        if started_ts is not None and stat.st_mtime < started_ts:
                            continue
                    except OSError:
                        continue
                    if f.suffix.lower() == ".docx":
                        docx = True
                    elif f.suffix.lower() == ".pdf":
                        pdf = True
            except OSError:
                pass

        elapsed = 0.0
        if self.run.started:
            try:
                # started 是 ISO 字符串（也会写进运行记录），这里必须先转成 datetime
                started_dt = datetime.fromisoformat(self.run.started)
                elapsed = max(0.0, (datetime.now() - started_dt).total_seconds())
            except (TypeError, ValueError):
                elapsed = 0.0

        export_pdf = bool((self.opts or {}).get("export_pdf", True))
        complete = False
        if st == "idle" and (docx or pdf):
            # 服务重启后加载的历史任务：输出目录里已有报告，说明上次已经完成
            complete = True
            pct, stage = 100, "报告已完成"
        elif st == "done":
            complete = docx or pdf
            if complete:
                pct, stage = 100, "报告已完成"
            else:
                pct, stage = 98, "完成，但没找到报告文件"
        elif st == "paused":
            pct = min(95, max(self.run.pct, 5))
            reason = self.run.pause_reason or (self.paused or {}).get("reason", "")
            stage = f"已暂停：{reason}" if reason else "已暂停，等待用户操作"
        elif pdf:
            # PDF 只有模型交付 DOCX 后才会生成，因此它出现就代表报告已经完成
            complete = True
            pct, stage = 100, "报告已完成"
        elif st == "canceled":
            pct, stage = self.run.pct, "已中止"
        elif st == "failed":
            pct, stage = self.run.pct, "生成失败"
        elif docx:
            pct = 96
            if self.run.exit_code is None:
                stage = "正文已生成，等待 DSH 收尾"
            elif export_pdf:
                stage = "正文已生成，正在导出 PDF"
            else:
                stage = "正文已生成，正在收尾"
        elif st == "running":
            pct = 8 + 84 * min(1.0, elapsed / EXPECTED_SECONDS)
            pct = min(92, pct)
            if not self.run.last_output and elapsed < 20:
                stage = "启动 DSH"
            elif elapsed < 15:
                stage = "加载 skill · 读取材料"
            else:
                stage = "AI 正在撰写正文"
        else:
            pct, stage = 0, "未开始"

        pct = int(min(100, max(0, round(pct))))
        if complete:
            pct = max(pct, self.run.pct)
            self.run.pct = pct
        elif st == "running":
            # 写作期间只进不退，但没完成前最高只到 99%
            pct = max(pct, min(self.run.pct, 99))
            self.run.pct = pct
        else:
            # done/failed/canceled/idle：没有完成就不能显示 100%
            pct = min(pct, 99)

        return {
            "pct": pct,
            "stage": stage,
            "state": st,
            "elapsed": int(elapsed),
            "complete": complete,
            "export_pdf": export_pdf,
            "steps": self._progress_steps(st, docx, pdf, export_pdf, complete),
        }

    def _progress_steps(self, st: str, docx: bool, pdf: bool,
                        export_pdf: bool, complete: bool) -> list:
        """给前端渲染的步骤清单；每步只表达确定的状态。"""
        def state_of(done: bool, active: bool = False) -> str:
            if done:
                return "done"
            if active:
                return "active"
            if st == "failed":
                return "failed"
            if st == "canceled":
                return "canceled"
            if st == "paused":
                return "paused"
            return "pending"

        write_done = docx or complete or st == "done"
        pdf_done = pdf
        pdf_skipped = not export_pdf and not pdf
        if pdf_done:
            pdf_state = "done"
        elif pdf_skipped:
            pdf_state = "skipped"
        elif docx and st == "running":
            pdf_state = "active"
        elif st == "done" and docx and not pdf:
            pdf_state = "failed"
        else:
            pdf_state = state_of(False)

        if docx or complete:
            docx_state = "done"
        elif st == "done":
            docx_state = "failed"
        else:
            docx_state = state_of(False)

        finish_state = "done" if complete else state_of(False)
        if not complete and st == "done":
            finish_state = "failed"

        return [
            {"key": "start", "label": "启动 DSH",
             "state": "done" if (self.run.started or complete) else
                      state_of(False, st == "running")},
            {"key": "write", "label": "读取材料 · 撰写正文",
             "state": state_of(write_done, st == "running" and not docx)},
            {"key": "docx", "label": "生成 DOCX", "state": docx_state},
            {"key": "pdf", "label": "导出 PDF" if export_pdf else "PDF（未勾选）",
             "state": pdf_state},
            {"key": "finish", "label": "报告完成", "state": finish_state},
        ]

    def usage(self):
        """当前/最近一次运行的 token 统计。"""
        if self.run.state == "running" or self.run.usage.get("steps"):
            return self.run.usage
        return self.last_usage or self.run.usage

    def public(self):
        files = []
        if self.out_dir.exists():
            for f in sorted(self.out_dir.iterdir()):
                if f.is_file() and f.suffix.lower() in {".docx", ".pdf", ".md", ".png", ".jpg"}:
                    files.append({"name": f.name, "size": human_size(f.stat().st_size),
                                  "mtime": f.stat().st_mtime})
        return {
            "id": self.id,
            "info": self.info,
            "opts": self.opts,
            "created": self.created,
            "dir": str(self.path),
            "cover": str(self.cover_path) if self.cover_path.exists() else None,
            "materials": self.materials,
            "template": self.template,
            "files": files,
            "progress": self.progress(),
            "run": {
                "state": self.run.state,
                "started": self.run.started,
                "finished": self.run.finished,
                "exit": self.run.exit_code,
                "answer": self.run.answer,
                "usage": self.usage(),
                "pause_reason": self.run.pause_reason or (self.paused or {}).get("reason", ""),
            },
        }


TASK_INDEX: dict[str, Task] = {}
INDEX_LOCK = threading.Lock()


def new_task(info: dict) -> Task:
    # 目录名保持纯 ASCII；中文信息存在 task.json 里用于展示。
    # 工作区本身的 ACL（WRITE_DAC/WRITE_OWNER）由安装器或用户保证。
    stamp = "t" + datetime.now().strftime("%y%m%d-%H%M%S")
    base = TASKS / stamp
    path, n = base, 1
    while path.exists():
        n += 1
        path = TASKS / f"{stamp}-{n}"
    path.mkdir(parents=True, exist_ok=True)
    tid = path.name
    t = Task(tid, path, dict(info))
    t.init_dirs()
    t.save()
    with INDEX_LOCK:
        TASK_INDEX[tid] = t
    return t


def get_task(tid: str) -> Task | None:
    with INDEX_LOCK:
        t = TASK_INDEX.get(tid)
    if t:
        return t
    p = TASKS / tid
    if p.exists():
        t = Task.load(p)
        if t:
            with INDEX_LOCK:
                TASK_INDEX[tid] = t
            return t
    return None


def delete_task(t: Task) -> dict:
    """删掉一个任务：先停掉在跑的 dsh，再把整个任务目录连文件夹一起删。"""
    sid = t.run.session_id or (t.paused or {}).get("session_id") or ""
    if t.run.state == "running" and t.run.proc:
        kill_tree(t.run.proc.pid)
        t.run.state = "canceled"
    # 只允许删 TASKS 里面的目录
    target = t.path.resolve()
    root = TASKS.resolve()
    if target == root or root not in target.parents:
        return {"ok": False, "error": f"拒绝删除任务目录之外的路径：{target}"}
    try:
        shutil.rmtree(target)
    except Exception as e:
        log(f"删除任务目录失败：{target}\n{traceback.format_exc()}")
        return {"ok": False, "error": f"删除失败：{e}"}
    with INDEX_LOCK:
        TASK_INDEX.pop(t.id, None)
    if sid:
        threading.Thread(target=lambda: cleanup_dsh_session(sid), daemon=True).start()
    return {"ok": True, "deleted": str(target), "id": t.id}


# ---------------------------------------------------------------- 材料入库
def _copy_image(src: Path, t: Task, prefix: str) -> dict:
    t.img_dir.mkdir(parents=True, exist_ok=True)
    name = f"{prefix}_{safe_name(src.stem, 'img', 30)}{src.suffix.lower()}"
    dst = t.img_dir / name
    n = 1
    while dst.exists():
        n += 1
        dst = t.img_dir / f"{prefix}_{safe_name(src.stem, 'img', 30)}_{n}{src.suffix.lower()}"
    shutil.copy2(src, dst)
    try:
        from PIL import Image
        with Image.open(dst) as im:
            size = f"{im.width}×{im.height}"
            mode = im.mode
    except Exception:
        size, mode = "", ""
    rel = dst.relative_to(t.path).as_posix()
    return {"name": dst.name, "path": rel, "size": size, "mode": mode}


def collect_folder_files(folder: Path, warn: list) -> list[Path]:
    """递归收集文件夹里可入库的实验材料（图片/PDF/Word/文本）。"""
    found = []
    skipped = skipped_big = 0
    total = 0
    try:
        entries = sorted(folder.rglob("*"), key=lambda p: str(p).lower())
    except OSError as e:
        warn.append(f"{folder}：读取失败（{e}）")
        return []
    for p in entries:
        try:
            if not p.is_file():
                continue
            rel = p.relative_to(folder)
            if any(part.startswith(".") for part in rel.parts):
                continue
            if any(part in SKIP_DIR_NAMES for part in rel.parts[:-1]):
                continue
            if p.name.startswith("~$") or p.name.lower() in ("thumbs.db", "desktop.ini"):
                continue
            if p.suffix.lower() not in MATERIAL_EXT:
                skipped += 1
                continue
            size = p.stat().st_size
            if size > FOLDER_MAX_FILE_BYTES:
                skipped_big += 1
                continue
            if total + size > FOLDER_MAX_TOTAL_BYTES:
                warn.append(f"{folder}：内容超过 {FOLDER_MAX_TOTAL_BYTES // (1024 * 1024)} MB，已停止继续读取")
                break
            found.append(p)
            total += size
            if len(found) >= FOLDER_MAX_FILES:
                warn.append(f"{folder}：文件超过 {FOLDER_MAX_FILES} 个，只读取前 {FOLDER_MAX_FILES} 个")
                break
        except OSError:
            continue
    if skipped:
        warn.append(f"{folder}：跳过 {skipped} 个不支持的文件")
    if skipped_big:
        warn.append(f"{folder}：跳过 {skipped_big} 个超过 {FOLDER_MAX_FILE_BYTES // (1024 * 1024)} MB 的文件")
    return found


def ingest_file(t: Task, src: Path, warn: list, display_name: str | None = None) -> dict:
    """把一份材料复制进任务目录并抽取文本/图片。"""
    src = Path(src)
    named = Path(display_name) if display_name else src
    ext = named.suffix.lower() or src.suffix.lower()
    kind = kind_of(ext)
    dst = t.mat_dir / safe_name(named.name, "材料" + ext, 80)
    n = 1
    while dst.exists():
        n += 1
        dst = t.mat_dir / f"{safe_name(named.stem, '材料', 60)}_{n}{ext}"
    shutil.copy2(src, dst)

    entry = {
        "name": dst.name,
        "path": dst.relative_to(t.path).as_posix(),
        "kind": kind,
        "size": human_size(dst.stat().st_size),
        "bytes": dst.stat().st_size,
        "note": "",
        "text": "",
        "pages": None,
        "images": [],
    }

    try:
        if ext in IMG_EXT:
            from PIL import Image
            with Image.open(dst) as im:
                entry["note"] = f"图片 {im.width}×{im.height}"
                entry["px"] = f"{im.width}×{im.height}"

        elif ext in TEXT_EXT:
            text = read_text_any(dst)
            # 跟随 Markdown / HTML 里的本地图片
            refs = MD_IMG_RE.findall(text) + HTML_IMG_RE.findall(text)
            picked = 0
            for ref in refs:
                p = _resolve_ref(ref, src.parent)
                if p and p.exists() and p.suffix.lower() in IMG_EXT:
                    entry["images"].append(_copy_image(p, t, safe_name(dst.stem, "材料", 20)))
                    picked += 1
            entry["note"] = f"{len(text)} 字" + (f"，关联图片 {picked} 张" if picked else "")
            txt = t.txt_dir / f"{dst.stem}.txt"
            txt.write_text(text, "utf-8")
            entry["text"] = txt.relative_to(t.path).as_posix()

        elif ext in PDF_EXT:
            text, pages = extract_pdf(dst)
            entry["pages"] = pages
            if len(re.sub(r"\s", "", text)) < 30 * max(1, pages or 1):
                entry["note"] = f"{pages} 页，未提取到文字（可能是扫描件），建议补图片"
                warn.append(f"{dst.name}：PDF 无文字层，可能是扫描件")
            else:
                entry["note"] = f"{pages} 页，已抽取 {len(text)} 字"
            txt = t.txt_dir / f"{dst.stem}.txt"
            txt.write_text(text, "utf-8")
            entry["text"] = txt.relative_to(t.path).as_posix()

        elif ext in DOCX_EXT:
            real = dst
            if ext == ".doc":
                conv = convert_doc_to_docx(dst)
                if conv:
                    real = conv
                else:
                    warn.append(f"{dst.name}：旧版 .doc 无法解析，建议另存为 .docx")
            if real.suffix.lower() == ".docx":
                text, imgs = extract_docx(real, t)
                entry["images"] = imgs
                entry["note"] = f"{len(text)} 字" + (f"，内嵌图片 {len(imgs)} 张" if imgs else "")
                txt = t.txt_dir / f"{dst.stem}.txt"
                txt.write_text(text, "utf-8")
                entry["text"] = txt.relative_to(t.path).as_posix()
        else:
            entry["note"] = "未识别的格式，仅作为附件交给 DSH"
    except Exception as e:
        entry["note"] = f"解析失败：{e}"
        warn.append(f"{dst.name}：解析失败（{e}）")
        log(traceback.format_exc())

    return entry


def _remove_template_files(t: Task):
    """删除任务里当前记录的模板文件（用于替换/移除模板）。"""
    old = t.template or {}
    root = t.path.resolve()
    for key in ("path", "text"):
        rel = old.get(key)
        if not rel:
            continue
        try:
            p = (t.path / rel).resolve()
            if root in p.parents and p.is_file():
                p.unlink()
        except OSError:
            pass


def _write_docx_template_outline(path: Path, out: Path) -> tuple[int, int]:
    """给 DSH 一份可快速浏览的模板大纲；正式排版仍以模板 DOCX 为准。"""
    import docx
    d = docx.Document(str(path))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    lines = ["# 实验报告模板大纲（辅助阅读；以模板 DOCX 本身为准）", ""]
    lines += paras[:400]
    if len(paras) > 400:
        lines.append(f"……其余 {len(paras) - 400} 段略")
    if d.tables:
        lines += ["", "## 表格"]
        for i, tb in enumerate(d.tables, 1):
            try:
                head = " | ".join(c.text.strip() for c in tb.rows[0].cells) if tb.rows else ""
            except Exception:
                head = ""
            lines.append(f"- 表 {i}：{len(tb.rows)} 行 × {len(tb.columns)} 列"
                         + (f"；首行：{head[:200]}" if head else ""))
    out.write_text("\n".join(lines), "utf-8")
    return len(paras), len(d.tables)


def ingest_template(t: Task, src: Path, display_name: str | None = None,
                    warn: list | None = None) -> dict:
    """把教师提供的报告模板复制进任务目录，并生成可读大纲。"""
    src = Path(src)
    named = Path(display_name) if display_name else src
    ext = (named.suffix or src.suffix).lower()
    if ext not in TEMPLATE_EXT:
        raise ValueError("模板只支持 DOCX / DOC / DOTX / PDF")
    t.template_dir.mkdir(parents=True, exist_ok=True)
    dst = t.template_dir / safe_name(named.name, "实验报告模板" + ext, 80)
    n = 1
    while dst.exists():
        n += 1
        dst = t.template_dir / f"{safe_name(named.stem, '实验报告模板', 60)}_{n}{ext}"
    shutil.copy2(src, dst)

    entry = {
        "name": dst.name,
        "path": dst.relative_to(t.path).as_posix(),
        "kind": "模板",
        "size": human_size(dst.stat().st_size),
        "bytes": dst.stat().st_size,
        "note": "",
        "text": "",
        "pages": None,
    }

    try:
        if ext in (".doc", ".dotx"):
            conv = convert_doc_to_docx(dst)
            if not conv:
                entry["note"] = "旧版/模板格式无法转换为 DOCX；DSH 可能无法读取"
                if warn is not None:
                    warn.append(f"{dst.name}：无法转换，建议另存为 .docx 后再上传")
                return entry
            target = dst.with_suffix(".docx")
            if target.exists():
                target = t.template_dir / f"{safe_name(dst.stem, '实验报告模板', 60)}_converted.docx"
            conv.replace(target)
            try:
                dst.unlink()
            except OSError:
                pass
            entry.update(name=target.name, path=target.relative_to(t.path).as_posix(),
                         kind="Word 模板", size=human_size(target.stat().st_size),
                         bytes=target.stat().st_size)
            dst = target

        if ext in (".docx", ".dotx", ".doc"):
            outline = t.template_dir / f"{dst.stem}_大纲.txt"
            paras, tables = _write_docx_template_outline(dst, outline)
            entry["kind"] = "Word 模板"
            entry["note"] = f"{paras} 段，{tables} 个表格"
            entry["text"] = outline.relative_to(t.path).as_posix()
        elif ext == ".pdf":
            text, pages = extract_pdf(dst)
            entry["kind"] = "PDF 模板"
            entry["pages"] = pages
            if len(re.sub(r"\s", "", text)) < 30 * max(1, pages or 1):
                entry["note"] = f"{pages} 页，未提取到文字（可能是扫描件），只能参考版式"
                if warn is not None:
                    warn.append(f"{dst.name}：PDF 模板没有文字层，只能参考版式")
            else:
                entry["note"] = f"{pages} 页，已抽取 {len(text)} 字（仅作结构参考）"
            txt = t.template_dir / f"{dst.stem}_大纲.txt"
            txt.write_text(text, "utf-8")
            entry["text"] = txt.relative_to(t.path).as_posix()
    except Exception as e:
        entry["note"] = f"模板解析失败：{e}"
        if warn is not None:
            warn.append(f"{dst.name}：模板解析失败（{e}）")
        log(traceback.format_exc())

    return entry


def _resolve_ref(ref: str, base: Path) -> Path | None:
    ref = ref.strip().strip("<>").strip()
    if not ref or ref.startswith(("http://", "https://", "data:")):
        return None
    if ref.lower().startswith("file:///"):
        ref = ref[8:]
    try:
        p = Path(ref.replace("/", "\\")) if re.match(r"^[a-zA-Z]:[\\/]", ref) else base / ref
        return p
    except Exception:
        return None


def extract_pdf(path: Path):
    vendor = CONFIG.get("vendor")
    if vendor and Path(vendor).exists() and vendor not in sys.path:
        sys.path.insert(0, vendor)
    try:
        from pypdf import PdfReader
    except Exception:
        return "", 0
    reader = PdfReader(str(path))
    chunks = []
    for i, page in enumerate(reader.pages):
        try:
            chunks.append(page.extract_text() or "")
        except Exception:
            chunks.append("")
    return "\n".join(chunks).strip(), len(reader.pages)


def extract_docx(path: Path, t: Task):
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        for row in tb.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    imgs = []
    try:
        for rel in d.part.rels.values():
            if "image" in rel.reltype:
                part = rel.target_part
                ext = Path(part.partname).suffix or ".png"
                tmp = t.mat_dir / f"__tmp{ext}"
                tmp.write_bytes(part.blob)
                imgs.append(_copy_image(tmp, t, safe_name(path.stem, "docx", 20)))
                tmp.unlink(missing_ok=True)
    except Exception:
        log("docx 图片抽取失败\n" + traceback.format_exc())
    return "\n".join(parts), imgs


def convert_doc_to_docx(path: Path) -> Path | None:
    """用已安装的 Word 把 .doc 另存为 .docx（失败就算了）。"""
    out = path.with_suffix(".converted.docx")
    ps = (
        "$w=New-Object -ComObject Word.Application; $w.Visible=$false; "
        f"$d=$w.Documents.Open('{str(path)}',$false,$true); "
        f"$d.SaveAs2('{str(out)}',16); $d.Close($false); $w.Quit()"
    )
    try:
        subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                       timeout=120, capture_output=True)
    except Exception:
        return None
    return out if out.exists() else None


def write_manifest(t: Task):
    lines = ["# 材料清单", "", f"任务：{t.info.get('course','')} / {t.info.get('title','')}",
             f"生成时间：{datetime.now():%Y-%m-%d %H:%M}", "",
             "| # | 文件 | 类型 | 大小 | 抽取内容 | 说明 |", "|---|---|---|---|---|---|"]
    for i, m in enumerate(t.materials, 1):
        lines.append(f"| {i} | {m['path']} | {m['kind']} | {m.get('size','')} | "
                     f"{m.get('text') or '—'} | {m.get('note','')} |")
        for img in m.get("images") or []:
            lines.append(f"| | 　↳ {img['path']} | 图片 | | | {img.get('size','')} |")
    if t.template:
        lines += ["", "## 实验报告模板（教师提供）",
                  f"- `{t.template.get('path') or '—'}`（{t.template.get('kind','模板')}，"
                  f"{t.template.get('size','')}）{t.template.get('note','')}"]
        if t.template.get("text"):
            lines.append(f"- 模板大纲：`{t.template['text']}`")
        lines.append("- 模板只决定报告结构与版式，不是本次实验的事实来源。")
    lines += ["", "说明：表格中的“抽取内容”是软件自动生成的可读文本副本，图片在 `材料/附件图/` 下。",
              "所有实验事实必须回到原始材料核对。"]
    (t.mat_dir / "材料清单.md").write_text("\n".join(lines), "utf-8")


# ---------------------------------------------------------------- 提示词
WIN_BAD = '\\/:*?"<>|'


def clean_outname(name: str) -> str:
    """把用户填的文件名洗成安全的主文件名（不含扩展名）。空串表示由 AI 自己决定。"""
    s = (name or "").strip()
    if not s:
        return ""
    for ch in WIN_BAD + "\r\n\t":
        s = s.replace(ch, " ")
    s = re.sub(r"\s+", " ", s).strip().strip(".")
    # 去掉常见扩展名，避免出现 报告.docx.docx
    for ext in (".docx", ".doc", ".pdf", ".md"):
        if s.lower().endswith(ext):
            s = s[: -len(ext)].strip()
            break
    # 去掉首尾的点与空格（Windows 不允许以点结尾）
    s = s.strip(" .")
    return s[:80]


def out_docx(t: Task, opts: dict) -> Path:
    """主交付 DOCX 的目标路径：填了文件名就用它，否则沿用默认的 实验报告.docx。"""
    name = clean_outname((opts or {}).get("outname") or "")
    return t.out_dir / (f"{name}.docx" if name else "实验报告.docx")


def build_prompt(t: Task, opts: dict) -> str:
    info = t.info
    course = info.get("course", "").strip()
    title = info.get("title", "").strip()
    skill_root = Path(CONFIG["skill_root"])
    tpl = t.template or None
    perm = permission_mode(opts)
    if perm == "danger-full-access":
        perm_block = """- 本次运行权限：**完全访问**（DSH 命令沙箱和审批已关闭）。
  如果材料或用户要求使用 MATLAB / Octave 等外部程序，可以实际运行并保留真实输出/图；
  仍然禁止编造结果，启动失败要如实记录。不要访问 HTTP 接口或操控 ReportWriter。
"""
    elif perm == "read-only":
        perm_block = "- 本次运行权限：只读。不要尝试写文件；这会阻碍报告交付。\n"
    else:
        perm_block = ""

    cur = (opts.get("course_key") or "").strip()
    if cur and cur != "auto":
        c = next((x for x in COURSES if x["key"] == cur), None)
    elif cur == "none":
        c = None
    else:
        c = match_course(f"{course} {title}")

    if c:
        res = (
            f"- 本次课程资源：**{c['name']}**\n"
            f"- 课程提示词（先读）：`{skill_root / c['prompt']}`\n"
            f"- 报告模板 DOCX（排版依据）：`{skill_root / c['template']}`\n"
        )
        if c.get("examples") and (skill_root / c["examples"]).exists():
            res += f"- 完成版实例目录（仅参照版式/语气，禁止照搬数据与结论）：`{skill_root / c['examples']}`\n"
    else:
        res = (
            f"- 课程「{course or '未填写'}」当前**未在本 skill 登记**。\n"
            f"- 请按 `{skill_root / 'references' / 'general-stem-lab-prompt.md'}` 和 "
            f"`{skill_root / 'references' / 'common-rules.md'}` 建立符合理工科实验报告逻辑的结构，"
            f"不要套用单片机/数电/模电等其他课程的固定章节与专业规则。\n"
            f"- 报告排版以本任务自带封面为基准（A4、页边距上下 2.54cm 左右 3.17cm、"
            f"正文中文宋体、西文 Times New Roman、一级标题黑体）。\n"
        )
    if tpl:
        res += "- **本任务已上传教师模板：模板的结构与版式优先于上述课程模板。**\n"

    if tpl:
        template_block = f"""## 二、实验报告模板（教师提供，优先）
- 模板文件：`{tpl.get('path')}`（{tpl.get('kind') or '模板'}，{tpl.get('size') or ''}）
- 模板说明：{tpl.get('note') or '—'}
"""
        if tpl.get("text"):
            template_block += f"- 模板大纲/文本副本：`{tpl['text']}`（快速浏览用；实际以模板文件为准）\n"
        template_block += """
**模板使用规则（硬要求）：**
1. 动笔前先读取模板与大纲，梳理章节顺序、标题层级、表格字段、评分项、图表/编号格式。
2. **教师模板优先于 skill 的课程模板**：结构、章节、表格字段、版式冲突时以教师模板为准；
   实验事实、专业规则与真实性约束仍按 skill 和材料执行。
3. 主交付仍以封面 DOCX 为基底：保留第 1 页封面与分节符，从第 2 页开始按模板组织正文；
   可从模板复制样式、表格、页眉页脚和页面设置，但不要改动封面页和分节符。
4. 模板若自带封面、姓名/学号栏或签署页：不要复制、不要填写；封面使用 ReportWriter 已生成的页面。
5. 模板中的占位符、提示语、示例数据必须替换为本次实验内容，**不得照抄模板示例数据**。
6. 模板要求但材料没有的数据不要编造；留空或写“材料未提供”，并在 `制作说明.md` 里列出。
7. PDF 模板只能沿用其章节、字段和文字要求，Word 样式按封面文档与通用排版规则处理。
"""
    else:
        template_block = """## 二、实验报告模板
- 本任务未上传教师模板。按「四、课程资源」中的课程模板与通用规则执行。
"""

    mat_lines = []
    for m in t.materials:
        bits = [m["kind"], m.get("size", "")]
        if m.get("pages"):
            bits.append(f"{m['pages']} 页")
        if m.get("note"):
            bits.append(m["note"])
        line = f"- `{m['path']}`（" + "，".join([b for b in bits if b]) + "）"
        if m.get("text"):
            line += f"；文本副本：`{m['text']}`"
        mat_lines.append(line)
        for img in (m.get("images") or [])[:20]:
            mat_lines.append(f"    - 关联图片：`{img['path']}`")
    if not mat_lines:
        mat_lines = ["- （未添加材料文件，请以用户在提示中粘贴的文字内容为准）"]

    paste = (opts.get("paste_text") or "").strip()
    paste_block = ""
    if paste:
        paste_file = t.mat_dir / "粘贴的实验内容.md"
        paste_file.write_text(paste, "utf-8")
        paste_block = (
            f"\n用户直接粘贴的实验内容（同样已保存到 `{paste_file.relative_to(t.path).as_posix()}`）：\n"
            "```\n" + paste[:6000] + "\n```\n"
        )

    extra = (opts.get("extra") or "").strip()
    extra_block = f"\n## 七、教师/用户附加要求（优先）\n{extra}\n" if extra else ""

    timeout_min = max(5, int(int(opts.get("timeout") or CONFIG["timeout"]) / 60))

    want = clean_outname(opts.get("outname") or "")
    main_docx = out_docx(t, opts)
    name_block = ""
    if want:
        name_block = (
            f"- **文件名（硬要求）**：主交付必须命名为 `{main_docx.name}`，一字不改，"
            f"不要加日期、姓名、v1 之类的后缀；`DONE:` 行里也必须是这个文件名。\n"
        )

    return f"""# 任务：撰写并交付本次实验报告

## 零、你的身份与边界（先读这一段）
- 你就是这份报告的**作者**，唯一目标：在本目录产出 `输出/{main_docx.name}`（PDF 可选）。
- 工作目录就是本任务根目录 `{t.path}`，材料 / 封面 / 输出都在这里，**在这里写文件是允许的**。
- 不要访问任何 HTTP 接口、不要启动或操控「ReportWriter」、不要找本地服务、不要试图提权或在别的目录里找捷径——
  那些都不是你的任务，只会浪费整轮时间。
- 某个工具报错（例如 pwsh 报沙箱 ACL 拒绝）时，**换一条路**继续，不要在同一处反复试探：
  用 write/read 工具直接落盘，或用 `{CONFIG['python']}` 跑 python-docx；只要最终 DOCX 出来就算成功。
  关键错误原文记进 `制作说明.md` 即可。
- **暂停规则**：只有确实缺少用户操作、材料或信息、无法继续时，才停止本轮，并**只输出一行**
  `PAUSE: <需要用户做什么>`，不要输出 DONE。ReportWriter 会暂停，用户补充后继续；
  能根据现有材料完成的部分先完成，不要因为小的不确定就暂停。
{perm_block}
请**先加载 skill `nuist-stem-lab-report`** 并严格按它的工作流程执行（读材料 → 识别课程 → 判断模式 → 选资源 → 提取证据 → 生成 → 导出 → 验收 → 交付）。

## 一、封面（已由「ReportWriter」生成，不要重做）
- 封面文件：`{t.cover_path}`（A4 纵向；标题「{cover_title(course, title)}」26pt 黑体加粗居中；封面页之后已插入**分节符**）
- 实验课名称：{course}
- 实验名称：{title}

**隐私约束（硬要求）**：封面第 1 页里有学生的个人信息（姓名、班级、学号、任课教师、实验日期），
ReportWriter 刻意**没有**把它们写进本提示词，所以你不需要、也不可以去获取它们：
- 为了追加正文你可以打开封面 DOCX，但**不要 dump、引用、复述或记录**其中的姓名 / 班级 / 学号 / 任课教师 / 日期；
- 正文、图题、表注、页眉页脚、致谢、`制作说明.md` 里**一律不得出现**任何个人信息；
- 也不要写进文档属性（`core_properties` 的作者 / 最后修改者 / 备注等）；
- 正文提到封面时只写「见封面」，不要抄封面内容；也不要自己编造姓名学号。
- 教师模板里若含前任学生/示例姓名、学号、签名等个人信息，**不要复制**到正文、表格、页眉页脚或文档属性中。

要求：以该封面 DOCX 为基底，**在分节符之后追加正文**（正文从第 2 页开始），封面页的文字、字号、居中与分页一律不要改动。

{template_block}
## 三、实验内容材料（唯一的事实来源）
材料目录：`{t.mat_dir}`
材料清单：`{t.mat_dir / '材料清单.md'}`

> 模板只决定报告的结构与版式，不是本次实验的事实来源；模板里的示例数据、结论不得当成实测数据。

{chr(10).join(mat_lines)}
{paste_block}
读取方式：DOCX/PDF 用 `{CONFIG['python']}` + python-docx / pypdf 取文本；PNG/JPG 用 `read_image` 直接看。**先逐个读完材料再动笔**。

## 四、课程资源
{res}
## 五、输出要求
- 输出目录：`{t.out_dir}`
- 主交付：`{t.out_dir / '实验报告.docx'}`（封面 + 正文；正文结构、编号、图表题注按教师模板（如有）、skill 与课程资源）
{name_block}
- **PDF 不用你做**：不要调用 Word COM / win32com / docx2pdf / soffice 去转 PDF。
  本机的 DSH 沙箱里 Word 起不来（报「Word 无法创建工作文件，请检查临时环境变量」，随后卡死十几分钟），
  这一坑已经踩过，**不要在 Word/PowerShell COM 上浪费任何时间**；软件会在你交付 DOCX 后自己转 PDF。
- 过程笔记写到 `{t.out_dir / '制作说明.md'}`：材料清单、缺失证据、待用户补充项、版本与命名。**不要把内部提示、占位符、待补说明塞进正式报告正文。**
- 完成后逐项自检：分页（封面独立第 1 页）、图表编号连续、正文引用一致、无空白页、图片清晰度、字体字号统一。
  **自检用 python-docx 复读 DOCX 即可，不要打开 Word。**
- 最后一行输出：`DONE: <实验报告 docx 的绝对路径>`

## 六、真实性约束
- 所有步骤、命令、截图、数据、结论都必须能在材料中找到依据；材料里没有的实测值、运行结果、截图**一律不编造**，不要用“预计”“可能”之类的字眼伪装成实测。
- 实验性质（实物/仿真/结合）材料里说不清时，先按材料能支撑的写法完成，并在 `制作说明.md` 里列出待确认项。
- 不联网检索；需要理论依据时读 skill 的课程提示词与公开常识。
- 时间预算约 {timeout_min} 分钟，请先保证交付一份完整可读的 DOCX，再考虑润色。
{extra_block}
现在开始。"""


def build_resume_prompt(t: Task, opts: dict, reason: str) -> str:
    """暂停后继续：沿用原 DSH 会话，只补充新增材料和继续要求。"""
    main_docx = out_docx(t, opts)
    mat_lines = []
    for m in t.materials:
        bits = [m.get("kind", ""), m.get("size", "")]
        if m.get("note"):
            bits.append(m["note"])
        line = f"- `{m['path']}`（" + "，".join(b for b in bits if b) + "）"
        if m.get("text"):
            line += f"；文本副本：`{m['text']}`"
        mat_lines.append(line)
    if not mat_lines:
        mat_lines = ["- （当前没有材料文件）"]

    paste = (opts.get("paste_text") or "").strip()
    paste_block = ""
    if paste:
        paste_file = t.mat_dir / "继续补充的实验内容.md"
        paste_file.write_text(paste, "utf-8")
        paste_block = (
            f"\n用户在继续前补充的文字（已保存到 `{paste_file.relative_to(t.path).as_posix()}`）：\n"
            "```\n" + paste[:6000] + "\n```\n"
        )

    extra = (opts.get("extra") or "").strip()
    extra_block = f"\n## 四、继续前新增要求\n{extra}\n" if extra else ""
    tpl_line = ""
    if t.template:
        tpl_line = f"- 教师模板仍有效：`{t.template.get('path')}`（结构与版式继续优先）\n"

    return f"""# 继续实验报告任务

用户已经暂停本轮，并完成了需要手动操作或补充材料的步骤。请从当前进度继续，**不要重新开始**。

## 一、上次暂停原因
{reason or '用户手动暂停'}

## 二、当前材料（可能已有新增）
材料目录：`{t.mat_dir}`
材料清单：`{t.mat_dir / '材料清单.md'}`
{tpl_line}{chr(10).join(mat_lines)}
{paste_block}
先重新读取新增/变化的材料，再继续写。原有 `任务说明.md` 里的真实性、隐私和封面约束全部继续有效。

## 三、继续要求
- **不要重做**已经完成且正确的章节；只补完缺失内容并统一图表编号。
- 如果仍然缺少用户操作/材料，继续输出一行 `PAUSE: <需要用户做什么>`，不要输出 DONE。
- 否则完成报告，最后一行输出：`DONE: {(t.out_dir / main_docx.name).resolve()}`
- 主交付：`{t.out_dir / main_docx.name}`（封面仍保留，正文从第 2 页开始）
- 真实记录进度：不要编造新增材料里没有的数据；新增材料不足时写进 `制作说明.md`。
{extra_block}
现在继续。"""


# ---------------------------------------------------------------- 运行 dsh
def start_run(t: Task, opts: dict, resume: bool = False) -> dict:
    if t.run.state == "running":
        return {"ok": False, "error": "任务已在运行中"}

    perm = permission_mode(opts)
    resume_session = ""
    base_usage = None
    if resume:
        if t.run.state != "paused":
            return {"ok": False, "error": "当前任务没有处于暂停状态"}
        paused = t.paused or {}
        resume_session = str(paused.get("session_id") or t.run.session_id or "").strip()
        if not resume_session:
            return {"ok": False, "error": "没有可继续的 DSH 会话，请重新开始"}
        base_usage = dict(t.run.usage)
        prompt = build_resume_prompt(t, opts, str(paused.get("reason") or "用户手动暂停"))
        prompt_path = t.path / "继续说明.md"
    else:
        prompt = build_prompt(t, opts)
        prompt_path = t.path / "任务说明.md"
    prompt_path.write_text(prompt, "utf-8")

    if not t.cover_path.exists():
        build_cover_docx(t.info, t.cover_path)

    t.run = RunState()
    if resume:
        t.run.session_id = resume_session
        t.run.usage = base_usage or dict(t.run.usage)
    t.run.state = "running"
    t.run.started = datetime.now().isoformat(timespec="seconds")
    t.run.prompt_path = str(prompt_path)
    # 日志/回答里若出现封面上的个人信息，落盘前替换成 *，避免日志文件外泄
    t.run.redact = [(t.info.get(k) or "").strip()
                    for k in ("name", "cls", "sid", "teacher")
                    if len((t.info.get(k) or "").strip()) >= 2]
    t.run.add("app", f"任务目录：{t.path}")
    t.run.add("app", f"提示词：{prompt_path.name}（{len(prompt)} 字）")
    if t.template:
        t.run.add("app", f"教师模板：{t.template.get('path')}（优先于课程模板）")
    t.run.add("app", f"执行权限：{PERMISSION_LABELS.get(perm, perm)}")
    if perm == "danger-full-access":
        t.run.add("app", "⚠ 已启用完全访问：DSH 命令沙箱和审批关闭，模型可执行任意本机命令")
    if resume:
        t.run.add("app", f"继续 DSH 会话：{resume_session}")
    t.run.add("app", "正在启动 DSH …")

    dsh = CONFIG["dsh"]
    if not Path(dsh).exists() and not shutil.which(dsh):
        t.run.state = "failed"
        t.run.add("error", f"找不到 dsh 命令：{dsh}，请在 config.json 里修改 dsh 路径")
        return {"ok": False, "error": "dsh 不存在"}

    api_key = str(opts.get("api_key") or "").strip()
    cmd, auth_note = dsh_launch_command(api_key, resume_session)
    mode = "api_key" if api_key else dsh_auth_state()[0]
    t.run.add("app", f"鉴权：{auth_note}")
    if mode == "missing":
        t.run.add("error", "未检测到 DSH 登录凭据或 DEEPSEEK_API_KEY；如果 DSH 报 no API key，"
                           "请先打开 DeepSeek Harness 登录账号，或在界面里填写 DEEPSEEK_API_KEY 后再试")
    elif mode == "account_stale":
        t.run.add("error", "DSH 账号凭据不可用（可能是 mock/开发凭据、过期或未完成登录）；"
                           "请打开 DeepSeek Harness 重新登录，或在界面里填写 DEEPSEEK_API_KEY")

    # 只保留 DSH_HOME，剥掉 DSH_SESSION_ID / DSH_PROFILE / DSH_SHELL 等
    # 会话级变量：否则子进程会误以为自己是"桌面会话里的嵌套 shell"，
    # 继承到不属于它的沙箱与会话状态。
    env = {k: v for k, v in os.environ.items()
           if not k.startswith("DSH_") or k == "DSH_HOME"}
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    if api_key:
        # 只传给本次 DSH 子进程，不落盘、不写日志
        env["DEEPSEEK_API_KEY"] = api_key
    # 父进程的 DSH_* 已被过滤掉；这里按本轮选择显式设置
    env["DSH_PERMISSION_MODE"] = perm

    try:
        proc = subprocess.Popen(
            cmd,
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            cwd=str(t.path), env=env,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except Exception as e:
        t.run.state = "failed"
        t.run.add("error", f"启动失败：{e}")
        return {"ok": False, "error": str(e)}

    t.run.proc = proc
    t.run.add("app", f"已启动 dsh headless（pid {proc.pid}），把材料交给模型…")

    def feed():
        try:
            proc.stdin.write(prompt.encode("utf-8"))
            proc.stdin.close()
        except Exception as e:
            t.run.add("error", f"写入提示词失败：{e}")

    def emit(name, line):
        if name == "stdout":
            try:
                ev = json.loads(line)
            except Exception:
                t.run.add("stdout", line)
                t.run.answer += redact_text(line, t.run.redact) + "\n"
            else:
                t.run.handle_event(ev)
        else:
            t.run.add(name, line)

    def pump(stream, name, sink):
        buf = b""
        try:
            while True:
                chunk = stream.read(1)
                if not chunk:
                    break
                buf += chunk
                if chunk in b"\n":
                    emit(name, decode_bytes(buf))
                    buf = b""
            if buf:
                emit(name, decode_bytes(buf))
        except Exception:
            pass
        finally:
            try:
                stream.close()
            except Exception:
                pass

    threading.Thread(target=feed, daemon=True).start()
    out_thread = threading.Thread(target=pump, args=(proc.stdout, "stdout", None), daemon=True)
    err_thread = threading.Thread(target=pump, args=(proc.stderr, "stderr", None), daemon=True)
    out_thread.start()
    err_thread.start()

    def drain_pumps():
        # 进程退出后等 stdout/stderr 读完，确保 final 和最后一个 step 的 token 用量都已入库
        for th in (out_thread, err_thread):
            th.join(5)

    def schedule_session_cleanup():
        if not CONFIG.get("cleanup_dsh_session", True):
            return
        sid = t.run.session_id
        if not sid:
            return
        run = t.run

        def work():
            if cleanup_dsh_session(sid):
                log(f"已清理 DSH headless 会话：{sid}")
                if t.run is run:
                    run.add("app", "已清理本次 DSH headless 会话（不会留在 DSH 未分组里）")
                    save_run_log(t)

        threading.Thread(target=work, daemon=True).start()

    def wait():
        timeout = int(opts.get("timeout") or CONFIG["timeout"])
        try:
            code = proc.wait(timeout=timeout)
        except subprocess.TimeoutExpired:
            kill_tree(proc.pid)
            drain_pumps()
            t.run.add("error", f"超过 {timeout // 60} 分钟仍未完成，已中止")
            t.run.state = "failed"
            t.paused = None
            t.run.finished = datetime.now().isoformat(timespec="seconds")
            t.run.exit_code = -1
            save_run_log(t)
            t.last_usage = dict(t.run.usage)
            t.save()
            schedule_session_cleanup()
            return
        drain_pumps()
        t.run.exit_code = code
        t.run.finished = datetime.now().isoformat(timespec="seconds")
        if t.run.cancel_requested:
            t.run.state = "canceled"
            t.paused = None
            save_run_log(t)
            t.last_usage = dict(t.run.usage)
            t.save()
            schedule_session_cleanup()
            return
        dm = re.search(r"DONE:\s*(.+?\.docx)", t.run.answer)
        has_docx = any(t.out_dir.glob("*.docx"))
        pause_mark = re.search(r"(?m)^PAUSE:\s*(.+?)\s*$", t.run.answer)
        reason = ""
        if t.run.pause_requested and not (code == 0 and (dm or has_docx)):
            reason = "用户手动暂停"
        elif pause_mark:
            reason = pause_mark.group(1).strip() or "模型请求用户操作"
        if reason:
            if not t.run.session_id:
                t.run.state = "failed"
                t.paused = None
                t.run.add("error", "暂停失败：DSH 会话尚未建立，无法继续")
                save_run_log(t)
                t.last_usage = dict(t.run.usage)
                t.save()
                return
            t.run.state = "paused"
            t.run.pause_reason = reason
            t.run.add("app", f"⏸ 已暂停：{reason}")
            t.paused = {
                "session_id": t.run.session_id,
                "reason": reason,
                "usage": dict(t.run.usage),
                "pct": t.run.pct,
                "at": datetime.now().isoformat(timespec="seconds"),
            }
            t.last_usage = dict(t.run.usage)
            save_run_log(t)
            t.save()
            return
        want_docx = out_docx(t, opts)
        want_name = clean_outname((opts or {}).get("outname") or "")
        if code == 0 and (dm or has_docx):
            t.run.add("app", "✅ 报告已生成" + (f"：{dm.group(1).strip()}" if dm else ""))
            if want_name:
                if want_docx.is_file():
                    t.run.add("app", f"文件名已按你的要求设为 `{want_docx.name}`")
                else:
                    t.run.add("error", f"注意：没找到你指定的 `{want_docx.name}`，下面按实际产出的文件交付")
            if opts.get("export_pdf", True):
                target = None
                # 用户指定了文件名时，优先用它
                if want_name and want_docx.is_file():
                    target = want_docx
                if target is None and dm:
                    cand = Path(dm.group(1).strip().strip("`"))
                    if cand.is_file():
                        target = cand
                if target is None:
                    found = sorted(t.out_dir.glob("*.docx"), key=lambda p: p.stat().st_mtime)
                    target = found[-1] if found else None
                if target is not None:
                    t.run.add("app", "正在用本机 Word 转 PDF（首次启动可能较慢）…")
                    res = export_pdf(target)
                    if res.get("ok"):
                        t.run.add("app", f"📄 PDF 已生成：{res['pdf']}")
                    else:
                        t.run.add("error", f"PDF 转换失败：{res.get('error')}")
            # 等 PDF 导出结束再标记完成：前端会一直轮询到此时刻，100% 才代表整份交付完成
            if t.run.state == "running":
                t.run.state = "done"
            t.paused = None
            t.run.pause_reason = ""
        else:
            t.run.state = "failed"
            t.paused = None
            t.run.add("error", f"DSH 退出码 {code}，未确认产出 DOCX")
            lines, _ = t.run.snapshot(0)
            auth_errors = ("no api key", "missing_credential",
                           "account_sign_in_required", "account_token_invalid")
            if any(any(k in txt.lower() for k in auth_errors) for _, _, txt in lines):
                t.run.add("error", "DSH 鉴权失败：请打开 DeepSeek Harness 重新登录账号，"
                                   "或在界面里填写 DEEPSEEK_API_KEY 后重试")
        save_run_log(t)
        t.last_usage = dict(t.run.usage)
        t.save()
        schedule_session_cleanup()

    threading.Thread(target=wait, daemon=True).start()
    return {"ok": True, "taskId": t.id}


def decode_bytes(b: bytes) -> str:
    for enc in ("utf-8", "gbk"):
        try:
            return b.decode(enc)
        except Exception:
            continue
    return b.decode("utf-8", "replace")


def kill_tree(pid: int):
    try:
        subprocess.run(["taskkill", "/F", "/T", "/PID", str(pid)],
                       capture_output=True, timeout=30)
    except Exception:
        pass


def save_run_log(t: Task):
    lines, _ = t.run.snapshot(0)
    body = "\n".join(f"[{s}] {txt}" for _, s, txt in lines)
    u = t.run.usage
    if u.get("steps"):
        body += (f"\n[usage] input={u.get('inputTokens', 0)} output={u.get('outputTokens', 0)} "
                 f"cache_read={u.get('cacheReadTokens', 0)} cache_write={u.get('cacheWriteTokens', 0)} "
                 f"total={u.get('totalTokens', 0)} steps={u.get('steps', 0)}")
    try:
        (t.path / "运行记录.log").write_text(body, "utf-8")
    except Exception:
        pass


def export_pdf(docx_path: Path) -> dict:
    """用本机 Word 把 DOCX 转 PDF。路径一律转绝对路径：Word 的 cwd 是 system32，相对路径必失败。"""
    docx_path = Path(docx_path).resolve()
    out = docx_path.with_suffix(".pdf")
    # 关闭隐藏对话框、拼写/语法检查和链接更新，尽量缩短 Word COM 的启动/打开时间。
    ps = (
        "$ErrorActionPreference='Stop';"
        "$w=New-Object -ComObject Word.Application;"
        "$w.Visible=$false; $w.DisplayAlerts=0; $w.ScreenUpdating=$false;"
        "$w.Options.UpdateLinksAtOpen=$false;"
        "$w.Options.CheckSpellingAsYouType=$false;"
        "$w.Options.CheckGrammarAsYouType=$false;"
        "$w.Options.SaveInterval=0;"
        f"$d=$w.Documents.Open('{docx_path}',$false,$true,$false);"
        f"$d.ExportAsFixedFormat('{out}',17,$false);"
        "$d.Close($false); $w.Quit()"
    )
    started = time.time()
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
            capture_output=True, timeout=180,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        elapsed = time.time() - started
        if out.exists():
            log(f"PDF 导出完成：{out}（{elapsed:.1f} 秒）")
            return {"ok": True, "pdf": str(out)}
        log(f"PDF 导出失败（{elapsed:.1f} 秒）：{decode_bytes(r.stderr)[-400:]}")
        return {"ok": False, "error": decode_bytes(r.stderr)[-400:] or "转换失败（未安装 Word？）"}
    except Exception as e:
        log(f"PDF 导出异常：{e}")
        return {"ok": False, "error": str(e)}


# ---------------------------------------------------------------- HTTP
class Handler(BaseHTTPRequestHandler):
    server_version = "ReportWriter/1.0"

    def log_message(self, fmt, *args):
        pass

    # ---------- 基础 ----------
    def _send(self, code, body: bytes, ctype="application/json; charset=utf-8", extra=None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        try:
            self.wfile.write(body)
        except Exception:
            pass

    def _json(self, obj, code=200):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode("utf-8"))

    def _err(self, msg, code=400):
        self._json({"ok": False, "error": msg}, code)

    def _body(self) -> bytes:
        n = int(self.headers.get("Content-Length") or 0)
        return self.rfile.read(n) if n else b""

    def _json_body(self) -> dict:
        raw = self._body()
        if not raw:
            return {}
        try:
            return json.loads(raw.decode("utf-8"))
        except Exception:
            return {}

    def do_GET(self):
        try:
            self.route_get()
        except Exception as e:
            log(traceback.format_exc())
            self._err(f"服务端错误：{e}", 500)

    def do_POST(self):
        try:
            self.route_post()
        except Exception as e:
            log(traceback.format_exc())
            self._err(f"服务端错误：{e}", 500)

    # ---------- GET ----------
    def route_get(self):
        u = urlparse(self.path)
        p, q = unquote(u.path), parse_qs(u.query)

        if p == "/api/bootstrap":
            courses = [{"key": c["key"], "name": c["name"]} for c in COURSES]
            auth_mode, auth_note = dsh_auth_state()
            return self._json({
                "ok": True,
                "auth": {"mode": auth_mode, "note": auth_note},
                "dsh_permission_mode": permission_mode(),
                "cover_history": load_cover_history(),
                "courses": courses,
                "skill_root": CONFIG["skill_root"],
                "dsh": CONFIG["dsh"],
                "python": CONFIG["python"],
                "timeout": CONFIG["timeout"],
                "port": self.server.server_address[1],
                "today": datetime.now().strftime("%Y-%m-%d"),
                "school_text": CONFIG.get("school_text") or "南京信息工程大学",
                "school_font": CONFIG.get("school_font") or "毛体",
                "school_color": CONFIG.get("school_color") or "#C41E23",
            })

        if p == "/api/skill/update":
            with _SKILL_UPDATE_LOCK:
                state = dict(_SKILL_UPDATE)
            state["ok"] = True
            return self._json(state)

        if p == "/api/tasks":
            items = []
            for d in sorted(TASKS.iterdir(), key=lambda x: x.stat().st_mtime, reverse=True)[:30]:
                if not d.is_dir():
                    continue
                t = Task.load(d)
                if t:
                    items.append({"id": t.id, "info": t.info, "created": t.created})
            return self._json({"ok": True, "tasks": items})

        m = re.match(r"^/api/task/([^/]+)$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            return self._json({"ok": True, "task": t.public()})

        m = re.match(r"^/api/task/([^/]+)/log$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            start = int((q.get("from") or ["0"])[0])
            lines, total = t.run.snapshot(start)
            return self._json({
                "ok": True, "from": start, "total": total,
                "lines": lines,
                "state": t.run.state,
                "exit": t.run.exit_code,
                "files": t.public()["files"],
                "progress": t.progress(),
                "usage": t.usage(),
                "pause_reason": t.run.pause_reason or (t.paused or {}).get("reason", ""),
            })

        if p == "/api/cover.png":
            info = {
                "course": (q.get("course") or [""])[0],
                "title": (q.get("title") or [""])[0],
                "name": (q.get("name") or [""])[0],
                "cls": (q.get("cls") or [""])[0],
                "sid": (q.get("sid") or [""])[0],
                "teacher": (q.get("teacher") or [""])[0],
                "date": (q.get("date") or [""])[0],
                "school": (q.get("school") or [""])[0],
                "school_font": (q.get("school_font") or [""])[0],
                "school_color": (q.get("school_color") or [""])[0],
            }
            return self._send(200, render_cover_png(info), "image/png")

        if p.startswith("/api/file/"):
            # /api/file/<taskid>/<relpath>
            rest = p[len("/api/file/"):]
            if "/" in rest:
                tid, rel = rest.split("/", 1)
                t = get_task(tid)
                if t:
                    fp = (t.path / rel).resolve()
                    if str(fp).startswith(str(t.path.resolve())) and fp.exists():
                        ctype = mimetypes.guess_type(str(fp))[0] or "application/octet-stream"
                        return self._send(200, fp.read_bytes(), ctype)

        return self.static(p)

    def static(self, p: str):
        if p in ("/", ""):
            p = "/index.html"
        fp = (WEB / p.lstrip("/")).resolve()
        if not str(fp).startswith(str(WEB.resolve())) or not fp.exists() or not fp.is_file():
            return self._send(404, b"not found", "text/plain; charset=utf-8")
        ctype = mimetypes.guess_type(str(fp))[0] or "application/octet-stream"
        if ctype.startswith("text/") or ctype.endswith("javascript"):
            ctype += "; charset=utf-8"
        return self._send(200, fp.read_bytes(), ctype)

    # ---------- POST ----------
    def route_post(self):
        u = urlparse(self.path)
        p = unquote(u.path)

        if p == "/api/preview":
            info = self._json_body()
            return self._send(200, render_cover_png(info), "image/png")

        if p == "/api/cover":
            payload = self._json_body()
            info = clean_info(payload.get("info") or {})
            tid = payload.get("taskId")
            t = get_task(tid) if tid else None
            if t:
                t.info = info
                t.save()
            else:
                t = new_task(info)
            build_cover_docx(t.info, t.cover_path)
            history = save_cover_history(t.info)
            return self._json({"ok": True, "task": t.public(), "cover_history": history})

        if p == "/api/cover/history/clear":
            clear_cover_history()
            return self._json({"ok": True, "cover_history": []})

        if p == "/api/skill/update/ack":
            ack_skill_update()
            with _SKILL_UPDATE_LOCK:
                state = dict(_SKILL_UPDATE)
            state["ok"] = True
            return self._json(state)

        if p == "/api/shutdown":
            threading.Thread(target=lambda: (time.sleep(0.4), os._exit(0)), daemon=True).start()
            return self._json({"ok": True})

        m = re.match(r"^/api/task/([^/]+)/upload$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            raw = self._body()
            ctype = self.headers.get("Content-Type") or ""
            files = parse_multipart(raw, ctype)
            if not files:
                return self._err("没有收到文件")
            warn, added = [], []
            for filename, data in files:
                tmp = t.path / f"__upload_{uuid.uuid4().hex}{Path(filename).suffix.lower()}"
                tmp.write_bytes(data)
                try:
                    added.append(ingest_file(t, tmp, warn, display_name=filename))
                finally:
                    tmp.unlink(missing_ok=True)
            t.materials.extend(added)
            t.save()
            write_manifest(t)
            return self._json({"ok": True, "added": added, "warnings": warn, "task": t.public()})

        m = re.match(r"^/api/task/([^/]+)/template$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            if t.run.state == "running":
                return self._err("任务正在运行中，暂时不能更换模板")
            files = parse_multipart(self._body(), self.headers.get("Content-Type") or "")
            if not files:
                return self._err("没有收到模板文件")
            if len(files) > 1:
                return self._err("一次只能上传一个模板文件")
            filename, data = files[0]
            ext = Path(filename).suffix.lower()
            if ext not in TEMPLATE_EXT:
                return self._err("模板只支持 DOCX / DOC / DOTX / PDF")
            if len(data) > TEMPLATE_MAX_BYTES:
                return self._err("模板文件不能超过 30 MB")
            tmp = t.path / f"__template_{uuid.uuid4().hex}{ext}"
            tmp.write_bytes(data)
            warn = []
            try:
                tpl = ingest_template(t, tmp, display_name=filename, warn=warn)
            except Exception as e:
                log(traceback.format_exc())
                return self._err(f"模板解析失败：{e}")
            finally:
                tmp.unlink(missing_ok=True)
            _remove_template_files(t)
            t.template = tpl
            t.save()
            write_manifest(t)
            return self._json({"ok": True, "template": tpl, "warnings": warn, "task": t.public()})

        m = re.match(r"^/api/task/([^/]+)/template/remove$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            if t.run.state == "running":
                return self._err("任务正在运行中，暂时不能更换模板")
            _remove_template_files(t)
            t.template = None
            t.save()
            write_manifest(t)
            return self._json({"ok": True, "task": t.public()})

        m = re.match(r"^/api/task/([^/]+)/materials/paths$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            paths = (self._json_body().get("paths") or [])
            warn, added = [], []
            for ps in paths:
                raw = str(ps).strip().strip('"')
                src = Path(os.path.expandvars(os.path.expanduser(raw)))
                if src.is_dir():
                    files = collect_folder_files(src, warn)
                    if not files:
                        warn.append(f"{raw}：没有找到可读取的图片 / PDF / DOCX / 文本")
                    for f in files:
                        added.append(ingest_file(t, f, warn))
                elif src.exists() and src.is_file():
                    added.append(ingest_file(t, src, warn))
                else:
                    warn.append(f"{raw}：路径不存在")
            t.materials.extend(added)
            t.save()
            write_manifest(t)
            return self._json({"ok": True, "added": added, "warnings": warn, "task": t.public()})

        m = re.match(r"^/api/task/([^/]+)/materials/remove$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            name = self._json_body().get("path")
            t.materials = [x for x in t.materials if x["path"] != name]
            t.save()
            write_manifest(t)
            return self._json({"ok": True, "task": t.public()})

        m = re.match(r"^/api/task/([^/]+)/run$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            opts = self._json_body()
            if not t.materials and not (opts.get("paste_text") or "").strip():
                return self._err("还没有实验材料：请拖入文件或在文本框里粘贴实验内容")
            if t.run.state == "running":
                return self._err("任务正在运行中")
            if t.run.state == "paused":
                return self._err("任务已暂停，请使用「继续写报告」或先结束当前任务")
            opts["outname"] = clean_outname(opts.get("outname") or "")
            if not opts["outname"]:
                opts.pop("outname", None)
            # 记住本轮选项，历史任务里能还原
            t.opts = {
                "outname": opts.get("outname") or "",
                "export_pdf": bool(opts.get("export_pdf", True)),
                "course_key": opts.get("course_key") or "auto",
                "permission_mode": permission_mode(opts),
                "extra": opts.get("extra") or "",
            }
            t.save()
            res = start_run(t, opts)
            return self._json(res)

        m = re.match(r"^/api/task/([^/]+)/pause$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            if t.run.state != "running":
                return self._err("任务没有在运行")
            if not t.run.session_id:
                return self._err("DSH 会话还在启动中，请稍后再试")
            t.run.pause_requested = True
            t.run.add("app", "收到暂停请求，正在保存会话…")
            if t.run.proc:
                kill_tree(t.run.proc.pid)
            return self._json({"ok": True})

        m = re.match(r"^/api/task/([^/]+)/resume$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            if t.run.state != "paused":
                return self._err("当前任务没有暂停")
            opts = self._json_body()
            if not opts.get("permission_mode"):
                opts["permission_mode"] = (t.opts or {}).get("permission_mode") or permission_mode()
            if "export_pdf" not in opts:
                opts["export_pdf"] = (t.opts or {}).get("export_pdf", True)
            if not opts.get("course_key"):
                opts["course_key"] = (t.opts or {}).get("course_key") or "auto"
            if not opts.get("extra"):
                opts["extra"] = (t.opts or {}).get("extra") or ""
            opts["outname"] = clean_outname(opts.get("outname") or (t.opts or {}).get("outname") or "")
            if not opts["outname"]:
                opts.pop("outname", None)
            t.opts = {
                "outname": opts.get("outname") or "",
                "export_pdf": bool(opts.get("export_pdf", True)),
                "course_key": opts.get("course_key") or "auto",
                "permission_mode": permission_mode(opts),
                "extra": opts.get("extra") or "",
            }
            t.save()
            return self._json(start_run(t, opts, resume=True))

        m = re.match(r"^/api/task/([^/]+)/cancel$", p)
        if m:
            t = get_task(m.group(1))
            if not t or t.run.state not in ("running", "paused"):
                return self._err("没有正在运行或暂停的任务")
            sid = t.run.session_id or (t.paused or {}).get("session_id") or ""
            if t.run.state == "running" and t.run.proc:
                t.run.cancel_requested = True
                kill_tree(t.run.proc.pid)
            t.run.state = "canceled"
            t.run.pause_reason = ""
            t.paused = None
            t.run.add("app", "已手动中止")
            save_run_log(t)
            t.save()
            if sid:
                threading.Thread(target=lambda: cleanup_dsh_session(sid), daemon=True).start()
            return self._json({"ok": True})

        m = re.match(r"^/api/task/([^/]+)/delete$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            return self._json(delete_task(t))

        m = re.match(r"^/api/task/([^/]+)/open$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            target = (self._json_body().get("target") or "dir")
            path = {
                "dir": t.path,
                "mat": t.mat_dir,
                "out": t.out_dir,
                "cover": t.cover_path,
            }.get(target)
            if target in ("docx", "pdf"):
                cands = [t.out_dir / "实验报告" / ("实验报告.docx" if target == "docx" else "实验报告.pdf")]
                cands += list(t.out_dir.rglob(f"*.{target}"))
                path = next((c for c in cands if c.exists()), None)
            if not path or not Path(path).exists():
                return self._err("文件还不存在")
            os.startfile(str(path))
            return self._json({"ok": True})

        m = re.match(r"^/api/task/([^/]+)/pdf$", p)
        if m:
            t = get_task(m.group(1))
            if not t:
                return self._err("任务不存在", 404)
            docxs = [t.out_dir / "实验报告.docx"] + list(t.out_dir.rglob("*.docx"))
            src = next((d for d in docxs if d.exists()), None)
            if not src:
                return self._err("还没有 DOCX")
            return self._json(export_pdf(src))

        return self._err("未知接口", 404)


def clean_info(info: dict) -> dict:
    out = {}
    for k in ("course", "title", "name", "cls", "sid", "teacher", "date"):
        out[k] = str(info.get(k) or "").strip()
    # 封面校名样式：文字、字体、颜色（留空表示用 config.json 的默认值）
    out["school"] = str(info.get("school") or "").strip()[:60]
    font = str(info.get("school_font") or "").strip()
    out["school_font"] = font if font in SCHOOL_FONTS else ""
    rgb = parse_color(info.get("school_color"))
    out["school_color"] = ("#%02X%02X%02X" % rgb) if rgb else ""
    return out


def parse_multipart(body: bytes, ctype: str):
    m = re.search(r'boundary="?([^";]+)"?', ctype)
    if not m or not body:
        return []
    boundary = m.group(1).encode()
    out = []
    for part in body.split(b"--" + boundary):
        if not part.strip() or part.strip() == b"--":
            continue
        head, _, data = part.partition(b"\r\n\r\n")
        if not _:
            continue
        hd = head.decode("utf-8", "replace")
        fn = re.search(r'filename="([^"]*)"', hd)
        if not fn:
            continue
        name = fn.group(1)
        if not name:
            continue
        if data.endswith(b"\r\n"):
            data = data[:-2]
        # 浏览器可能带路径
        name = name.replace("\\", "/").split("/")[-1]
        out.append((name, data))
    return out


# ---------------------------------------------------------------- 启动
def free_port(start: int) -> int:
    for port in range(start, start + 40):
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
            s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            try:
                s.bind(("127.0.0.1", port))
                return port
            except OSError:
                continue
    raise RuntimeError("找不到可用端口")


def open_window(url: str):
    edge_paths = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ]
    profile = Path(os.environ.get("LOCALAPPDATA", str(ROOT))) / "ReportWriter" / "edge-profile"
    for e in edge_paths:
        if Path(e).exists():
            try:
                subprocess.Popen([e, f"--app={url}", "--window-size=1440,1000",
                                  f"--user-data-dir={profile}"])
                return
            except Exception:
                continue
    webbrowser.open(url)


def main():
    load_config()
    start_skill_update_check()
    port = free_port(int(CONFIG.get("port") or 8791))
    srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    url = f"http://127.0.0.1:{port}/"
    log(f"启动 http://127.0.0.1:{port}")
    print(url, flush=True)
    if CONFIG.get("open_browser", True):
        threading.Timer(0.8, lambda: open_window(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
